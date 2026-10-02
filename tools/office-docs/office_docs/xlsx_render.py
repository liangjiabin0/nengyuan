from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from .fonts import resolve_font_files


THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
ALT_FILL = PatternFill("solid", fgColor="F4F7FB")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NUMBER_FORMATS = {
    "number": "0.00",
    "int": "#,##0",
    "percent": "0.00%",
    "currency": '¥#,##0.00;¥-#,##0.00;"-"',
    "date": "YYYY-MM-DD",
    "text": "@",
}


def _font_name(preset: str) -> str:
    files = resolve_font_files()
    if preset == "data-export":
        return "Arial" if files.get("arial") else "Calibri"
    return "微软雅黑" if files.get("yahei") else "宋体"


def _width_for(title: str, rows: list[dict[str, Any]], key: str, given: Any) -> float:
    if given:
        return float(given)
    longest = len(str(title))
    for row in rows[:200]:
        longest = max(longest, len(str(row.get(key, ""))))
    return min(max(longest + 2, 10), 42)


def _expand_formula(template: str, columns: list[dict], row_idx: int) -> str:
    formula = template
    formula = formula.replace("{row}", str(row_idx))
    for index, col in enumerate(columns, start=1):
        letter = get_column_letter(index)
        formula = formula.replace("{col:" + col["key"] + "}", f"{letter}{row_idx}")
        formula = formula.replace("{" + col["key"] + "}", f"{letter}{row_idx}")
    return formula


def _add_chart(ws, sheet: dict, columns: list[dict], n_rows: int):
    for chart_spec in sheet.get("charts") or []:
        kind = str(chart_spec.get("kind") or "bar")
        title = str(chart_spec.get("title") or sheet.get("name") or "图表")
        y_keys = chart_spec.get("y") or []
        if isinstance(y_keys, str):
            y_keys = [y_keys]
        keys = [col["key"] for col in columns]
        if not y_keys:
            y_keys = [col["key"] for col in columns[1:3]]
        x_key = chart_spec.get("x") or (columns[0]["key"] if columns else None)
        if not x_key or x_key not in keys or n_rows < 1:
            continue
        x_idx = keys.index(x_key) + 1
        cats = Reference(ws, min_col=x_idx, min_row=2, max_row=1 + n_rows)
        if kind == "pie":
            chart = PieChart()
            y_idx = keys.index(y_keys[0]) + 1 if y_keys[0] in keys else 2
            data = Reference(ws, min_col=y_idx, min_row=1, max_row=1 + n_rows)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
        else:
            chart = LineChart() if kind == "line" else BarChart()
            for key in y_keys:
                if key not in keys:
                    continue
                y_idx = keys.index(key) + 1
                data = Reference(ws, min_col=y_idx, min_row=1, max_row=1 + n_rows)
                chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.y_axis.title = None
        chart.title = title
        chart.style = 10
        chart.y_axis.axId = 100
        anchor = str(chart_spec.get("anchor") or get_column_letter(len(columns) + 2) + "2")
        chart.width = 15
        chart.height = 8
        ws.add_chart(chart, anchor)


def render_xlsx(spec: dict, output: str | Path) -> Path:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    font_name = _font_name(spec.get("preset") or "cn-workbook")
    header_font = Font(name=font_name, bold=True, color="FFFFFF", size=11)
    body_font = Font(name=font_name, size=10.5, color="1F2328")
    note_font = Font(name=font_name, size=9, italic=True, color="57606A")

    first = True
    for sheet in spec.get("sheets") or []:
        ws = default if first else wb.create_sheet()
        first = False
        ws.title = sheet["name"]
        columns = sheet.get("columns") or []
        rows = sheet.get("rows") or []
        ws.sheet_properties.tabColor = "1F4E79"
        ws.page_setup.orientation = "landscape" if len(columns) > 6 else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75)
        ws.oddHeader.left.text = spec.get("title") or sheet["name"]
        ws.oddFooter.center.text = "&P / &N"
        ws.print_title_rows = "1:1"

        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(1, col_idx, col["title"])
            cell.font = header_font
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN
            ws.column_dimensions[get_column_letter(col_idx)].width = _width_for(col["title"], rows, col["key"], col.get("width"))

        ws.row_dimensions[1].height = 22
        ws.auto_filter.ref = f"A1:{get_column_letter(max(len(columns), 1))}{max(len(rows) + 1, 1)}"
        freeze = sheet.get("freeze") or "A2"
        if isinstance(freeze, str):
            ws.freeze_panes = freeze

        for r_idx, row in enumerate(rows, start=2):
            for c_idx, col in enumerate(columns, start=1):
                cell = ws.cell(r_idx, c_idx)
                formula = col.get("formula") or ""
                if formula:
                    cell.value = _expand_formula(formula, columns, r_idx)
                else:
                    cell.value = row.get(col["key"], "")
                cell.font = body_font
                cell.border = THIN
                cell.alignment = Alignment(
                    horizontal=col.get("align") or ("right" if col.get("format") in NUMBER_FORMATS and col.get("format") != "text" else "left"),
                    vertical="center",
                    wrap_text=True,
                )
                if r_idx % 2 == 0:
                    cell.fill = ALT_FILL
                fmt_key = str(col.get("format") or "")
                if fmt_key in NUMBER_FORMATS and fmt_key != "text":
                    cell.number_format = NUMBER_FORMATS[fmt_key]

        totals = sheet.get("totals") or {}
        sum_keys = totals.get("sum") or []
        if sum_keys:
            total_row = len(rows) + 2
            ws.cell(total_row, 1, "合计").font = Font(name=font_name, bold=True, size=10.5)
            keys = [col["key"] for col in columns]
            for key in sum_keys:
                if key not in keys:
                    continue
                idx = keys.index(key) + 1
                letter = get_column_letter(idx)
                cell = ws.cell(total_row, idx, f"=SUM({letter}2:{letter}{len(rows)+1})")
                cell.font = Font(name=font_name, bold=True, size=10.5)
                cell.number_format = NUMBER_FORMATS.get(columns[idx - 1].get("format") or "", "0.00")

        if sheet.get("notes"):
            note_row = len(rows) + (3 if sum_keys else 2)
            ws.cell(note_row, 1, sheet["notes"]).font = note_font
            if columns:
                ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(columns))

        _add_chart(ws, sheet, columns, len(rows))
        ws.sheet_view.showGridLines = False

    if spec.get("author"):
        wb.properties.creator = spec["author"]
    wb.properties.title = spec.get("title") or ""
    wb.save(str(output))
    return output
