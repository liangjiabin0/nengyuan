from __future__ import annotations

import json
import os
import shutil
import subprocess
import zipfile
from pathlib import Path


def _issues(items: list[str]) -> dict:
    return {"ok": not items, "issues": items}


def inspect_docx(path: str | Path) -> dict:
    path = Path(path)
    issues = []
    if not path.exists():
        return _issues([f"文件不存在：{path}"])
    if path.stat().st_size < 2000:
        issues.append("docx 过小，可能是空文档")
    try:
        from docx import Document
        doc = Document(str(path))
    except Exception as exc:
        return _issues([f"无法打开 docx：{exc}"])
    text = "\n".join(p.text for p in doc.paragraphs)
    if len(text.strip()) < 10:
        issues.append("正文几乎为空")
    cal = 0
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            name = (run.font.name or "") + str(run._element.xml)
            if "Calibri" in name and run.text.strip():
                cal += 1
    if cal > 8:
        issues.append("大量 Calibri 正文，中文报告应使用宋体/黑体/雅黑")
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        if "w:color w:val=\"0563C1\"" in xml or "themeColor=\"accent1\"" in xml:
            issues.append("标题可能仍是主题蓝色，应改为黑色")
    except Exception:
        pass
    return {
        "ok": not issues,
        "issues": issues,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "chars": len(text),
    }


def inspect_xlsx(path: str | Path) -> dict:
    path = Path(path)
    issues = []
    if not path.exists():
        return _issues([f"文件不存在：{path}"])
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), data_only=False)
    except Exception as exc:
        return _issues([f"无法打开 xlsx：{exc}"])
    sheets = []
    for ws in wb.worksheets:
        used = ws.max_row or 0
        cols = ws.max_column or 0
        if used < 2 or cols < 1:
            issues.append(f"{ws.title} 几乎没有数据")
        if ws.freeze_panes is None:
            issues.append(f"{ws.title} 未冻结表头")
        sheets.append({"name": ws.title, "rows": used, "cols": cols, "freeze": str(ws.freeze_panes)})
    return {"ok": not issues, "issues": issues, "sheets": sheets}


def inspect_csv(path: str | Path) -> dict:
    path = Path(path)
    issues = []
    if not path.exists():
        return _issues([f"文件不存在：{path}"])
    raw = path.read_bytes()
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        issues.append("CSV 使用 UTF-16，Excel 中文场景请用 utf-8-sig")
    text = path.read_text(encoding="utf-8-sig")
    lines = [line for line in text.splitlines() if line.strip()]
    if len(lines) < 2:
        issues.append("CSV 缺少表头或数据行")
    return {"ok": not issues, "issues": issues, "lines": len(lines), "has_bom": raw.startswith(b"\xef\xbb\xbf")}


def inspect_pptx(path: str | Path) -> dict:
    path = Path(path)
    issues = []
    if not path.exists():
        return _issues([f"文件不存在：{path}"])
    if path.stat().st_size < 2000:
        issues.append("pptx 过小，可能是空文件")
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
    except Exception as exc:
        return _issues([f"无法打开 pptx：{exc}"])
    slides = list(prs.slides)
    if not slides:
        issues.append("没有幻灯片")
    ratio = (prs.slide_width or 1) / (prs.slide_height or 1)
    if abs(ratio - (16 / 9)) > 0.02:
        issues.append(f"画布不是 16:9（{ratio:.3f}）")
    chars = 0
    cal = 0
    for slide in slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    text = run.text or ""
                    chars += len(text)
                    name = (run.font.name or "") + str(run._r.xml)
                    if "Calibri" in name and text.strip():
                        cal += 1
    if chars < 8:
        issues.append("几乎没有正文")
    if cal > 12:
        issues.append("大量 Calibri 正文，中文幻灯片应使用宋体/黑体/雅黑")
    return {
        "ok": not issues,
        "issues": issues,
        "slides": len(slides),
        "chars": chars,
        "ratio": round(ratio, 4),
    }


def _pdftoppm() -> str | None:
    env = os.environ.get("PDFTOPPM_PATH")
    if env and Path(env).exists():
        return env
    return shutil.which("pdftoppm")


def inspect_pdf(path: str | Path, render_dir: str | Path | None = None) -> dict:
    path = Path(path)
    issues = []
    if not path.exists():
        return _issues([f"文件不存在：{path}"])
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = len(reader.pages)
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        return _issues([f"无法打开 pdf：{exc}"])
    if pages < 1:
        issues.append("PDF 没有页面")
    if len(text.strip()) < 10:
        issues.append("未能抽出正文，可能字体未嵌入或页面为空白")
    if "\ufffd" in text:
        issues.append("抽出文本含替换字符，可能中文字体未正确嵌入")
    rendered = []
    bin_path = _pdftoppm()
    if render_dir and bin_path:
        out_dir = Path(render_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        prefix = out_dir / "page"
        result = subprocess.run(
            [bin_path, "-png", "-r", "110", str(path), str(prefix)],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if result.returncode != 0:
            issues.append(f"pdftoppm 失败：{(result.stderr or result.stdout)[:200]}")
        else:
            rendered = sorted(str(item) for item in out_dir.glob("page*.png"))
    elif render_dir and not bin_path:
        issues.append("未找到 pdftoppm，跳过页面渲染。可安装 Poppler 或设置 PDFTOPPM_PATH")
    return {
        "ok": not issues,
        "issues": issues,
        "pages": pages,
        "chars": len(text),
        "rendered": rendered,
    }


def dumps(result: dict) -> str:
    return json.dumps(result, ensure_ascii=False, indent=2)
