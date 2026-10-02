"""Deprecated for chat agents: use the official docx skill. Kept so office-docs CLI still imports."""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT

from .fonts import preset_font_names
from .text import split_cjk_ascii


BLACK = RGBColor(0x00, 0x00, 0x00)


def set_run_font(run, *, east: str, ascii_font: str, size_pt: float | None = None, bold: bool | None = None):
    run.font.name = ascii_font
    run.font.color.rgb = BLACK
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:cs"), ascii_font)


def add_mixed_text(paragraph, text: str, *, east: str, ascii_font: str, size_pt: float, bold: bool = False):
    if not text:
        return
    for chunk, is_cjk in split_cjk_ascii(text):
        run = paragraph.add_run(chunk)
        set_run_font(run, east=east, ascii_font=ascii_font, size_pt=size_pt, bold=bold)


def _set_style_font(style, *, east: str, ascii_font: str, size_pt: float, bold: bool = False):
    font = style.font
    font.name = ascii_font
    font.size = Pt(size_pt)
    font.bold = bold
    font.color.rgb = BLACK
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:cs"), ascii_font)


def _add_page_field(paragraph, instr: str = "PAGE"):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_end)


def _set_cell_border(cell, **edges):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, attrs in edges.items():
        element = tcBorders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tcBorders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def _set_table_layout(table, *, autofit: bool, widths_cm: list[float] | None = None):
    table.autofit = autofit
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "autofit" if autofit else "fixed")
    if widths_cm:
        grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
        for grid_col, width in zip(grid_cols, widths_cm):
            grid_col.set(qn("w:w"), str(round(Cm(width).twips)))


def _set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_row_keep_together(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = trPr.find(qn("w:cantSplit"))
    if cant_split is None:
        trPr.append(OxmlElement("w:cantSplit"))


def _set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = trPr.find(qn("w:tblHeader"))
    if tbl_header is None:
        trPr.append(OxmlElement("w:tblHeader"))


def apply_three_line_table(table, header: bool = True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    spec = {
        "top": {"val": "single", "sz": "12", "space": "0", "color": "000000"},
        "left": {"val": "nil"},
        "bottom": {"val": "single", "sz": "12", "space": "0", "color": "000000"},
        "right": {"val": "nil"},
        "insideH": {"val": "nil"},
        "insideV": {"val": "nil"},
    }
    for edge, attrs in spec.items():
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))
    if header and table.rows:
        for cell in table.rows[0].cells:
            _set_cell_border(
                cell,
                top={"val": "single", "sz": "12", "space": "0", "color": "000000"},
                bottom={"val": "single", "sz": "8", "space": "0", "color": "000000"},
                left={"val": "nil"},
                right={"val": "nil"},
            )
        _set_repeat_header(table.rows[0])


def configure_document(doc: Document, spec: dict) -> dict[str, str]:
    fonts = preset_font_names(spec["preset"])
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    if spec["preset"] == "memo":
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    else:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    _set_style_font(normal, east=fonts["east"], ascii_font=fonts["ascii"], size_pt=12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    heading_sizes = {1: 16, 2: 14, 3: 12}
    heading_space = {1: (12, 6), 2: (6, 3), 3: (6, 3)}
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        _set_style_font(
            style,
            east=fonts["east_heading"],
            ascii_font=fonts["ascii"],
            size_pt=heading_sizes[level],
            bold=True,
        )
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(heading_space[level][0])
        style.paragraph_format.space_after = Pt(heading_space[level][1])
        style.paragraph_format.first_line_indent = Cm(0)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_mixed_text(hp, spec.get("header") or spec.get("title") or "", east=fonts["east"], ascii_font=fonts["ascii"], size_pt=9)
    if spec.get("confidential"):
        hp.add_run("  内部资料")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if spec.get("footer"):
        add_mixed_text(fp, spec["footer"] + "  ", east=fonts["east"], ascii_font=fonts["ascii"], size_pt=9)
    if spec.get("page_numbers", True):
        _add_page_field(fp)
    return fonts


def add_title_block(doc: Document, spec: dict, fonts: dict[str, str]):
    if not spec.get("title"):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Cm(0)
    add_mixed_text(p, spec["title"], east=fonts["title"], ascii_font=fonts["ascii"], size_pt=18, bold=True)
    if spec.get("subtitle"):
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.paragraph_format.first_line_indent = Cm(0)
        add_mixed_text(s, spec["subtitle"], east=fonts["east"], ascii_font=fonts["ascii"], size_pt=12)
    meta = "  ".join(item for item in [spec.get("org"), spec.get("author"), spec.get("date")] if item)
    if meta:
        m = doc.add_paragraph()
        m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        m.paragraph_format.space_after = Pt(18)
        m.paragraph_format.first_line_indent = Cm(0)
        add_mixed_text(m, meta, east=fonts["east"], ascii_font=fonts["ascii"], size_pt=10.5)


def add_toc(doc: Document, spec: dict, fonts: dict[str, str]):
    if not spec.get("toc"):
        return
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Cm(0)
    add_mixed_text(heading, "目录", east=fonts["east_heading"], ascii_font=fonts["ascii"], size_pt=16, bold=True)
    entries = []
    for section in spec["sections"]:
        if section.get("heading"):
            entries.append((section["level"], section["heading"]))
    for level, title in entries:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.74 * max(level - 1, 0))
        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        add_mixed_text(p, title + "\t", east=fonts["east"], ascii_font=fonts["ascii"], size_pt=12)
    doc.add_page_break()


def add_paragraph(doc: Document, text: str, fonts: dict[str, str], *, indent: bool = True, size: float = 12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    add_mixed_text(p, text, east=fonts["east"], ascii_font=fonts["ascii"], size_pt=size)
    return p


def add_list(doc: Document, items: list[str], fonts: dict[str, str], ordered: bool):
    for index, item in enumerate(items, start=1):
        marker = f"{index}. " if ordered else "（%s）" % index
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        add_mixed_text(p, marker + item, east=fonts["east"], ascii_font=fonts["ascii"], size_pt=12)


def add_table(doc: Document, block: dict, fonts: dict[str, str]):
    if block.get("title"):
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.keep_with_next = True
        add_mixed_text(cap, block["title"], east=fonts["east"], ascii_font=fonts["ascii"], size_pt=10.5, bold=True)
    headers = block.get("headers") or []
    rows = block.get("rows") or []
    cols = max(len(headers), max((len(row) for row in rows), default=1), 1)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    widths = [float(value) for value in (block.get("widths_cm") or [])]
    autofit = bool(block.get("autofit", not bool(widths)))
    repeat_header = bool(block.get("repeat_header", True))
    keep_rows = bool(block.get("keep_rows", True))
    alignments = block.get("align") or []
    _set_table_layout(table, autofit=autofit, widths_cm=widths)
    if widths:
        for idx, width in enumerate(widths[:cols]):
            table.columns[idx].width = Cm(width)
    if headers:
        for idx in range(cols):
            cell = table.cell(0, idx)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_mixed_text(p, str(headers[idx] if idx < len(headers) else ""), east=fonts["east"], ascii_font=fonts["ascii"], size_pt=10, bold=True)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            requested = str(alignments[c_idx] if c_idx < len(alignments) else "center").lower()
            p.alignment = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }.get(requested, WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            value = row[c_idx] if c_idx < len(row) else ""
            add_mixed_text(p, str(value), east=fonts["east"], ascii_font=fonts["ascii"], size_pt=9.5)
    for row in table.rows:
        if keep_rows:
            _set_row_keep_together(row)
        for c_idx, cell in enumerate(row.cells):
            _set_cell_margins(cell, top=70, start=90, bottom=70, end=90)
            if c_idx < len(widths):
                cell.width = Cm(widths[c_idx])
    if repeat_header and headers:
        _set_repeat_header(table.rows[0])
    if str(block.get("style") or "three-line") == "three-line":
        apply_three_line_table(table, header=bool(headers))
    return table


def add_figure(doc: Document, block: dict, fonts: dict[str, str]):
    path = block.get("path") or ""
    if not path or not Path(path).exists():
        add_paragraph(doc, f"（缺图：{path or '未提供路径'}）", fonts, indent=False, size=10.5)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    width = Cm(float(block.get("width_cm") or 14))
    run.add_picture(path, width=width)
    if block.get("caption"):
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Cm(0)
        add_mixed_text(cap, block["caption"], east=fonts["east"], ascii_font=fonts["ascii"], size_pt=10.5)


def _mathml_to_omml(mathml: str):
    candidates = [
        Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
        Path(r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL"),
    ]
    xsl_path = next((path for path in candidates if path.exists()), None)
    if xsl_path is None:
        raise RuntimeError("MathML equation rendering requires Microsoft Office MML2OMML.XSL")
    transform = etree.XSLT(etree.parse(str(xsl_path)))
    source = etree.fromstring(mathml.encode("utf-8"))
    result = transform(source)
    return parse_xml(etree.tostring(result.getroot(), encoding="unicode"))


def _append_equation_number(paragraph, number: str, fonts: dict[str, str]):
    if not number:
        return
    add_mixed_text(paragraph, f"    （{number}）", east=fonts["east"], ascii_font=fonts["ascii"], size_pt=12)


def add_equation(doc: Document, block: dict, fonts: dict[str, str]):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    number = str(block.get("number") or "")
    omml = str(block.get("omml") or "").strip()
    mathml = str(block.get("mathml") or "").strip()
    if omml:
        p._p.append(parse_xml(omml))
        _append_equation_number(p, number, fonts)
        return
    if mathml:
        p._p.append(_mathml_to_omml(mathml))
        _append_equation_number(p, number, fonts)
        return
    latex = str(block.get("latex") or "")
    display = latex if not number else f"{latex}    （{number}）"
    add_mixed_text(p, display, east=fonts["east"], ascii_font=fonts["ascii"], size_pt=12, bold=False)


def add_code(doc: Document, text: str, fonts: dict[str, str]):
    for line in (text or "").split("\n") or [""]:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(line)
        set_run_font(run, east=fonts["east"], ascii_font="Consolas", size_pt=10)


def render_blocks(doc: Document, blocks: list[dict], fonts: dict[str, str]):
    for block in blocks:
        kind = block.get("type")
        if kind == "paragraph":
            add_paragraph(doc, block.get("text") or "", fonts, indent=bool(block.get("indent", True)))
        elif kind == "list":
            add_list(doc, block.get("items") or [], fonts, bool(block.get("ordered")))
        elif kind == "table":
            add_table(doc, block, fonts)
        elif kind == "figure":
            add_figure(doc, block, fonts)
        elif kind == "equation":
            add_equation(doc, block, fonts)
        elif kind == "quote":
            add_paragraph(doc, block.get("text") or "", fonts, indent=False, size=10.5)
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "code":
            add_code(doc, block.get("text") or "", fonts)


def render_docx(spec: dict, output: str | Path) -> Path:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    fonts = configure_document(doc, spec)
    core = doc.core_properties
    core.title = spec.get("title") or ""
    core.author = spec.get("author") or "能元"
    core.comments = "Generated by AI4ES office-docs"
    add_title_block(doc, spec, fonts)
    add_toc(doc, spec, fonts)
    for section in spec.get("sections") or []:
        if section.get("heading"):
            heading = doc.add_heading(section["heading"], level=int(section.get("level") or 1))
            heading.paragraph_format.first_line_indent = Cm(0)
            for run in heading.runs:
                set_run_font(
                    run,
                    east=fonts["east_heading"],
                    ascii_font=fonts["ascii"],
                    size_pt={1: 16, 2: 14, 3: 12}.get(int(section.get("level") or 1), 12),
                    bold=True,
                )
        render_blocks(doc, section.get("blocks") or [], fonts)
    if spec.get("references"):
        heading = doc.add_heading("参考文献", level=1)
        heading.paragraph_format.first_line_indent = Cm(0)
        for index, item in enumerate(spec["references"], start=1):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(-0.74)
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.line_spacing = 1.5
            add_mixed_text(p, f"[{index}] {item}", east=fonts["east"], ascii_font=fonts["ascii"], size_pt=10.5)
    doc.save(str(output))
    return output
