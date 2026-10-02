from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

from .fonts import preset_font_names


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def _set_run_fonts(run, *, east: str, ascii_font: str, size_pt: float, bold: bool = False) -> None:
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = DARK
    r_pr = run._r.get_or_add_rPr()
    for tag, name in (("a:latin", ascii_font), ("a:ea", east), ("a:cs", ascii_font)):
        el = r_pr.find(qn(tag))
        if el is None:
            el = etree.SubElement(r_pr, qn(tag))
        el.set("typeface", name)


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _write_lines(tf, lines: list[str], *, east: str, ascii_font: str, size_pt: float, bold: bool = False, bullet: bool = False):
    tf.word_wrap = True
    tf.clear()
    if not lines:
        return
    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.level = 0
        paragraph.space_after = Pt(8 if bullet else 4)
        run = paragraph.add_run()
        run.text = line
        _set_run_fonts(run, east=east, ascii_font=ascii_font, size_pt=size_pt, bold=bold)


def render_pptx(spec: dict, out_path: str | Path) -> Path:
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    fonts = preset_font_names(spec.get("preset") or "cn-business")
    east = fonts["east"]
    heading = fonts["east_heading"]
    ascii_font = fonts["ascii"] if spec.get("preset") == "en-report" else "Times New Roman"
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]

    for index, item in enumerate(spec.get("slides") or [], start=1):
        slide = prs.slides.add_slide(blank)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()

        title_box = _add_textbox(slide, Inches(0.6), Inches(0.28), Inches(12.1), Inches(1.05))
        title_box.text_frame.word_wrap = True
        _write_lines(
            title_box.text_frame,
            [str(item.get("title") or f"第{index}页")],
            east=heading,
            ascii_font=ascii_font,
            size_pt=28,
            bold=True,
        )

        subtitle = str(item.get("subtitle") or "").strip()
        body_top = Inches(1.35)
        if subtitle:
            sub_box = _add_textbox(slide, Inches(0.6), Inches(1.22), Inches(12.1), Inches(0.45))
            _write_lines(
                sub_box.text_frame,
                [subtitle],
                east=east,
                ascii_font=ascii_font,
                size_pt=16,
            )
            body_top = Inches(1.7)

        bullets = [str(line).strip() for line in (item.get("bullets") or []) if str(line).strip()]
        figures = [fig for fig in (item.get("figures") or []) if fig.get("path")]
        body_width = Inches(7.4) if figures else Inches(12.1)
        if bullets:
            body_box = _add_textbox(slide, Inches(0.6), body_top, body_width, Inches(5.2))
            _write_lines(
                body_box.text_frame,
                bullets,
                east=east,
                ascii_font=ascii_font,
                size_pt=18,
                bullet=True,
            )

        if figures:
            fig_left = Inches(8.2)
            fig_width = Inches(4.5)
            slot = Inches(4.8) / max(len(figures), 1)
            for fig_index, figure in enumerate(figures):
                image = Path(figure["path"])
                if not image.is_file():
                    continue
                top = body_top + Emu(int(slot) * fig_index)
                try:
                    slide.shapes.add_picture(str(image), fig_left, top, width=fig_width)
                except Exception:
                    continue

        footer = _add_textbox(slide, Inches(0.6), Inches(7.05), Inches(11.4), Inches(0.32))
        footer_bits = [str(spec.get("title") or "").strip(), str(spec.get("author") or "").strip(), str(index)]
        _write_lines(
            footer.text_frame,
            ["  ·  ".join(part for part in footer_bits if part)],
            east=east,
            ascii_font=ascii_font,
            size_pt=11,
        )

        notes = str(item.get("notes") or "").strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.save(str(out))
    return out
