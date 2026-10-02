from pathlib import Path
from typing import Any

from pptx import Presentation

from ppt_voice_builder.text.normalizer import normalize_text


def _shape_text(shape: Any) -> str:
    if bool(getattr(shape, "has_text_frame", False)):
        return (shape.text_frame.text or "").strip()
    return ""


def _slide_title(slide: Any) -> str:
    title = slide.shapes.title
    return title.text.strip() if title is not None and hasattr(title, "text") else ""


def _slide_body(slide: Any) -> str:
    title = slide.shapes.title
    parts: list[str] = []
    for shape in slide.shapes:
        if title is not None and shape == title:
            continue
        text = _shape_text(shape)
        if text:
            parts.append(text)
        if bool(getattr(shape, "has_table", False)):
            for row in shape.table.rows:
                cells = [
                    (cell.text_frame.text or "").strip()
                    for cell in row.cells
                    if (cell.text_frame.text or "").strip()
                ]
                if cells:
                    parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def inspect_slides(source_pptx: Path) -> list[dict[str, Any]]:
    source_pptx = source_pptx.expanduser().resolve()
    if not source_pptx.is_file() or source_pptx.suffix.lower() != ".pptx":
        raise FileNotFoundError(f"找不到有效的 PPTX 文件：{source_pptx}")
    presentation = Presentation(str(source_pptx))
    result: list[dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        raw = ""
        if bool(getattr(slide, "has_notes_slide", False)):
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                raw = frame.text or ""
        normalized = normalize_text(raw)
        result.append(
            {
                "slideIndex": index,
                "slideId": int(slide.slide_id),
                "title": _slide_title(slide),
                "bodyText": _slide_body(slide),
                "rawNotes": raw,
                "notes": normalized,
                "hasNote": bool(normalized),
                "characterCount": len(normalized),
            }
        )
    return result
