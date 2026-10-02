import os
from pathlib import Path

from pptx import Presentation


def write_slide_notes(
    source_pptx: Path,
    output_pptx: Path,
    notes_by_index: dict[int, str],
) -> Path:
    source_pptx = source_pptx.expanduser().resolve()
    output_pptx = output_pptx.expanduser().resolve()
    if not source_pptx.is_file() or source_pptx.suffix.lower() != ".pptx":
        raise FileNotFoundError(f"找不到有效的 PPTX 文件：{source_pptx}")
    presentation = Presentation(str(source_pptx))
    for index, slide in enumerate(presentation.slides, start=1):
        if index not in notes_by_index:
            continue
        text = str(notes_by_index[index] or "")
        notes_slide = slide.notes_slide
        frame = notes_slide.notes_text_frame
        frame.text = text
    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_pptx.with_name(f"{output_pptx.stem}.tmp{os.getpid()}.pptx")
    try:
        presentation.save(str(temporary))
        os.replace(temporary, output_pptx)
    finally:
        temporary.unlink(missing_ok=True)
    return output_pptx
