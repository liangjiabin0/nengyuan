from contextlib import suppress
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from ppt_voice_builder.domain.errors import PowerPointError


def slide_image_name(slide_index: int) -> str:
    return f"slide-{slide_index:03d}.png"


def _emu_to_px(value: int, slide_width_emu: int, target_width: int) -> int:
    if slide_width_emu <= 0:
        return 0
    return max(1, round(int(value) * target_width / slide_width_emu))


def _load_font(size: int):
    from PIL import ImageFont

    for candidate in (
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\msyh.ttf",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\arial.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        if Path(candidate).is_file():
            try:
                return ImageFont.truetype(candidate, size)
            except Exception:
                continue
    return ImageFont.load_default()


def export_slides_fallback(source_pptx: Path, output_dir: Path, width: int = 1600) -> list[str]:
    from io import BytesIO

    from PIL import Image, ImageDraw

    presentation = Presentation(str(source_pptx))
    slide_width = int(presentation.slide_width or Emu(9144000))
    slide_height = int(presentation.slide_height or Emu(5143500))
    height = max(1, round(width * slide_height / slide_width))
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        canvas = Image.new("RGB", (width, height), (15, 23, 42))
        draw = ImageDraw.Draw(canvas)
        for shape in slide.shapes:
            left = _emu_to_px(int(shape.left or 0), slide_width, width)
            top = _emu_to_px(int(shape.top or 0), slide_height, height)
            box_w = _emu_to_px(int(shape.width or 0), slide_width, width)
            box_h = _emu_to_px(int(shape.height or 0), slide_height, height)
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                try:
                    picture = Image.open(BytesIO(shape.image.blob)).convert("RGB")
                    picture = picture.resize((max(1, box_w), max(1, box_h)))
                    canvas.paste(picture, (left, top))
                    continue
                except Exception:
                    pass
            text = ""
            if bool(getattr(shape, "has_text_frame", False)):
                text = (shape.text_frame.text or "").strip()
            if not text:
                continue
            font = _load_font(max(14, min(36, box_h // 6 or 18)))
            draw.multiline_text((left + 8, top + 6), text, fill=(248, 250, 252), font=font, spacing=6)
        path = output_dir / slide_image_name(index)
        canvas.save(path, "PNG")
        written.append(str(path))
    return written


def export_slides_with_powerpoint(source_pptx: Path, output_dir: Path, width: int = 1600) -> list[str]:
    source_pptx = source_pptx.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    app: Any = None
    presentation: Any = None
    initialized = False
    try:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        initialized = True
        app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = app.Presentations.Open(str(source_pptx), WithWindow=False, ReadOnly=True)
        slide_count = int(presentation.Slides.Count)
        slide_width = float(presentation.PageSetup.SlideWidth)
        slide_height = float(presentation.PageSetup.SlideHeight)
        height = max(1, round(width * slide_height / slide_width)) if slide_width else 900
        written: list[str] = []
        for index in range(1, slide_count + 1):
            dest = output_dir / slide_image_name(index)
            presentation.Slides(index).Export(str(dest), "PNG", width, height)
            written.append(str(dest))
        presentation.Close()
        presentation = None
        app.Quit()
        app = None
        return written
    except Exception as exc:
        raise PowerPointError(f"导出幻灯片预览失败：{exc}") from exc
    finally:
        if presentation is not None:
            with suppress(Exception):
                presentation.Close()
        if app is not None:
            with suppress(Exception):
                app.Quit()
        if initialized:
            with suppress(Exception):
                pythoncom.CoUninitialize()


def export_slide_images(source_pptx: Path, output_dir: Path, width: int = 1600) -> dict[str, Any]:
    source_pptx = source_pptx.expanduser().resolve()
    if not source_pptx.is_file():
        raise FileNotFoundError(f"找不到 PPTX：{source_pptx}")
    try:
        paths = export_slides_with_powerpoint(source_pptx, output_dir, width)
        return {"engine": "powerpoint", "count": len(paths), "paths": paths}
    except Exception:
        pass
    try:
        from ppt_voice_builder.ppt.linux_engine import export_slides_with_libreoffice

        paths = export_slides_with_libreoffice(source_pptx, output_dir, width)
        return {"engine": "libreoffice", "count": len(paths), "paths": paths}
    except Exception:
        paths = export_slides_fallback(source_pptx, output_dir, width)
        return {"engine": "fallback", "count": len(paths), "paths": paths}
