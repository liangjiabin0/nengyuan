from pathlib import Path
from typing import Any

from pptx import Presentation

from ppt_voice_builder.domain.models import SlideNote
from ppt_voice_builder.text.normalizer import normalize_text


def _slide_title(slide: Any) -> str:
    title = slide.shapes.title
    return title.text.strip() if title is not None and hasattr(title, "text") else ""


def read_slide_notes(source_pptx: Path) -> list[SlideNote]:
    source_pptx = source_pptx.expanduser().resolve()
    if not source_pptx.is_file() or source_pptx.suffix.lower() != ".pptx":
        raise FileNotFoundError(f"找不到有效的 PPTX 文件：{source_pptx}")
    try:
        presentation = Presentation(str(source_pptx))
    except Exception as exc:
        raise FileNotFoundError(
            f"无法打开 PPTX（文件可能正在写入或不是完整压缩包）：{source_pptx}"
        ) from exc
    result: list[SlideNote] = []
    for index, slide in enumerate(presentation.slides, start=1):
        raw = ""
        if bool(getattr(slide, "has_notes_slide", False)):
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                raw = frame.text or ""
        normalized = normalize_text(raw)
        result.append(
            SlideNote(
                index, int(slide.slide_id), _slide_title(slide), raw, normalized, bool(normalized)
            )
        )
    return result
