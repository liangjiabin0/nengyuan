"""Explicit CSV/Excel network templates, parsed as data without formula execution."""

import csv


def read_template(source):
    if source.suffix.lower() == ".csv":
        with source.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            if "network" not in workbook:
                raise ValueError("Excel 模板必须包含 network 工作表")
            sheet = workbook["network"]
            if sheet.max_row > 5000 or sheet.max_column > 30:
                raise ValueError("表格超过模板规模限制")
            records = []
            for cells in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("算例模板不接受公式，请提交明确数值")
                records.append([cell.value for cell in cells])
            rows = [dict(zip(records[0], values)) for values in records[1:]]
        finally:
            workbook.close()
    if not rows or len(rows) > 5000:
        raise ValueError("表格为空或超过 5000 行")
    network = dict(schemaVersion="static-network/v1", buses=[], lines=[], loads=[])

    def number(row, key):
        if row.get(key) in ["", None]:
            raise ValueError("缺少列数值：" + key)
        return float(row[key])

    for row in rows:
        kind = str(row.get("component", "")).strip()
        if not kind:
            continue
        if kind == "network":
            network["baseMva"] = number(row, "base_mva")
            network["slackBus"] = str(row.get("bus", "")).strip()
        elif kind == "bus":
            network["buses"].append(
                dict(id=str(row["id"]).strip(), voltageKv=number(row, "voltage_kv"))
            )
        elif kind == "line":
            network["lines"].append(
                dict(
                    id=str(row["id"]).strip(),
                    **{
                        "from": str(row["from_bus"]).strip(),
                        "to": str(row["to_bus"]).strip(),
                    },
                    lengthKm=number(row, "length_km"),
                    rOhmPerKm=number(row, "r_ohm_per_km"),
                    xOhmPerKm=number(row, "x_ohm_per_km"),
                    cNfPerKm=number(row, "c_nf_per_km"),
                    maxCurrentKa=number(row, "max_i_ka"),
                )
            )
        elif kind == "load":
            network["loads"].append(
                dict(
                    bus=str(row["bus"]).strip(),
                    pMw=number(row, "p_mw"),
                    qMvar=number(row, "q_mvar"),
                )
            )
        else:
            raise ValueError("模板含未支持元件类型：" + kind)
    if "baseMva" not in network or not network.get("slackBus"):
        raise ValueError("缺少 network 行的 base_mva 和平衡母线 bus")
    return network
