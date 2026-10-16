"""One immutable run per process; communicate with a real stdio MCP server."""

import argparse
import asyncio
import contextlib
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import sys
import time
import shutil
from datetime import datetime, timezone
from analysis import planning_network, planning_result

HERE = Path(__file__).resolve().parent


def save(path, value):
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
    )


def event(phase, detail):
    print(
        json.dumps({"phase": phase, "detail": detail}, ensure_ascii=False), flush=True
    )


def unpack(reply):
    if reply.is_error:
        raise RuntimeError("MCP tool returned isError")
    result = reply.structured_content
    if result is None:
        for block in reply.content:
            if getattr(block, "type", "") == "text":
                try:
                    result = json.loads(block.text)
                    break
                except ValueError:
                    continue
    if isinstance(result, dict) and set(result) == {"result"}:
        result = result["result"]
    if not isinstance(result, dict):
        raise ValueError("MCP response has no structured object")
    return result


def prepare_flow(spec, directory):
    import pandapower as pp
    import pandapower.networks as networks

    config = spec["parameters"]
    factories = {"ieee9": networks.case9, "ieee14": networks.case14}
    if spec["modelId"] == "uploaded":
        data = json.loads((directory / "source.json").read_text())
        if data.get("schemaVersion") != "static-network/v1":
            net = pp.from_json(str(directory / "source.json"))
            net.load.loc[:, ["p_mw", "q_mvar"]] *= config["loadFactor"]
            dest = directory / "network.json"
            pp.to_json(net, str(dest))
            return dest
        net = pp.create_empty_network(sn_mva=data["baseMva"])
        buses = {
            row["id"]: pp.create_bus(net, vn_kv=row["voltageKv"], name=row["id"])
            for row in data["buses"]
        }
        pp.create_ext_grid(net, bus=buses[data["slackBus"]], vm_pu=1.0)
        for row in data["lines"]:
            if (
                abs(
                    net.bus.at[buses[row["from"]], "vn_kv"]
                    - net.bus.at[buses[row["to"]], "vn_kv"]
                )
                > 1e-6
            ):
                raise ValueError(
                    "Different voltage levels require an explicit transformer; this line import does not infer one"
                )
            pp.create_line_from_parameters(
                net,
                buses[row["from"]],
                buses[row["to"]],
                length_km=row["lengthKm"],
                r_ohm_per_km=row["rOhmPerKm"],
                x_ohm_per_km=row["xOhmPerKm"],
                c_nf_per_km=row["cNfPerKm"],
                max_i_ka=row["maxCurrentKa"],
                name=row["id"],
            )
        for row in data["loads"]:
            pp.create_load(net, buses[row["bus"]], p_mw=row["pMw"], q_mvar=row["qMvar"])
    elif spec["modelId"] not in factories:
        raise ValueError("Unsupported static model; import must be validated first")
    else:
        net = factories[spec["modelId"]]()
    net.load.loc[:, ["p_mw", "q_mvar"]] *= config["loadFactor"]
    path = directory / "network.json"
    pp.to_json(net, str(path))
    return path


def prepare_dynamic(spec, directory):
    import andes
    from openpyxl import load_workbook

    config = spec["parameters"]
    if spec["modelId"] not in ["kundur-full", "uploaded"]:
        raise ValueError("A validated dynamic model is required")
    workbook = load_workbook(
        directory / "source.xlsx"
        if spec["modelId"] == "uploaded"
        else andes.get_case("kundur/kundur_full.xlsx")
    )
    if "Bus" not in workbook or not any(
        name in workbook for name in ["GENROU", "GENCLS", "GENSAL"]
    ):
        raise ValueError("Dynamic synchronous machine parameters are required")
    if config.get("analysis") != "small_signal":
        channels = (workbook["Bus"].max_row - 1) + 2 * sum(
            workbook[name].max_row - 1
            for name in ["GENROU", "GENCLS", "GENSAL"]
            if name in workbook
        )
        if channels * math.ceil(config["endAtS"] / config["stepS"]) > 2000000:
            raise ValueError(
                "动态输出超过当前任务预算，请缩短仿真范围或减少模型规模；未启动求解"
            )
    if "Fault" in workbook:
        del workbook["Fault"]
    if config.get("analysis") == "small_signal":
        path = directory / "network.xlsx"
        workbook.save(path)
        return path
    sheet = workbook.create_sheet("Fault")
    sheet.append(["idx", "u", "name", "bus", "tf", "tc", "xf", "rf"])
    sheet.append(
        [
            1,
            1,
            "Requested three-phase fault",
            config["busId"],
            config["faultAtS"],
            config["clearAtS"],
            0.0001,
            0,
        ]
    )
    path = directory / "network.xlsx"
    workbook.save(path)
    return path


def dynamic_result(spec, raw):
    c = spec["parameters"]
    if raw.get("status") != "success" or not raw.get("simulation", {}).get("success"):
        return {
            "solverStatus": "diverged",
            "criteriaStatus": "not_evaluated",
            "message": raw.get("message", "TDS failed"),
            "metrics": {},
            "tables": {},
            "series": [],
        }
    data = raw["series"]
    times = data["timeS"]
    voltage = data["busVoltagePu"].get(str(c["busId"]))
    if (
        voltage is None
        or len(times) < 3
        or len(voltage) != len(times)
        or times[-1] < c["endAtS"] - c["stepS"] * 2
    ):
        raise ValueError("Dynamic output is incomplete or observed bus is missing")
    if any(b <= a for a, b in zip(times, times[1:])):
        raise ValueError("Dynamic time index is not strictly increasing")
    recovery = None
    for index, t in enumerate(times):
        if t < c["clearAtS"] or times[-1] - t < c["holdS"]:
            continue
        if all(c["voltageMinPu"] <= v <= c["voltageMaxPu"] for v in voltage[index:]):
            recovery = t - c["clearAtS"]
            break
    speed_id, speed = next(iter(data["generatorSpeedPu"].items()))
    rows = [
        {"timeS": t, "voltagePu": v, "speedPu": w}
        for t, v, w in zip(times, voltage, speed)
    ]
    extra = []
    for group, label, unit in [
        ("generatorSpeedPu", "转速", "p.u."),
        ("generatorAngleRad", "功角", "rad"),
        ("busFrequencyPu", "频率", "p.u."),
    ]:
        for idx, values in data.get(group, {}).items():
            if len(values) != len(times):
                raise ValueError("动态变量采样长度不完整")
            if group == "generatorSpeedPu" and idx == speed_id:
                continue
            extra.append(
                dict(
                    label=idx + " " + label,
                    unit="Hz" if group == "busFrequencyPu" else unit,
                    xUnit="s",
                    points=[
                        [
                            t,
                            v * data["baseFrequencyHz"]
                            if group == "busFrequencyPu"
                            else v,
                        ]
                        for t, v in zip(times, values)
                    ],
                )
            )
    response = {
        "solverStatus": "completed",
        "criteriaStatus": "pass"
        if recovery is not None and recovery <= c["recoverWithinS"]
        else "violation",
        "metrics": {
            "minVoltagePu": min(voltage),
            "recoveryS": recovery,
            "speedPeakPu": max(speed),
        },
        "series": [
            {
                "label": "Bus " + str(c["busId"]) + " 电压",
                "unit": "p.u.",
                "xUnit": "s",
                "points": [[t, v] for t, v in zip(times, voltage)],
            },
            {
                "label": speed_id + " 转速",
                "unit": "p.u.",
                "xUnit": "s",
                "points": [[t, v] for t, v in zip(times, speed)],
            },
        ],
        "tables": {"dynamic": rows},
        "warnings": [
            "仅评价指定母线的电压恢复；转子转速不等于母线频率，不代表系统整体稳定结论。"
        ],
    }
    response["series"] += extra
    if not data.get("busFrequencyPu"):
        response["warnings"].append("算例未配置母线频率测量模型，未生成母线频率曲线。")
    return response


async def execute(spec, directory):
    started = time.time()
    canonical = spec.get("runSpec")
    if canonical:
        if (
            canonical["metadata"]["parameters"] != spec["parameters"]
            or canonical["solver"]["backend"] != spec["engine"]
        ):
            raise ValueError("RunSpec 与待执行参数不一致")
        spec["parameters"] = canonical["metadata"]["parameters"]
    event("preparing", "校验模型并准备原生输入")
    registered = json.loads((HERE / "capabilities.json").read_text())["engines"][
        spec["engine"]
    ]
    if importlib.metadata.version(spec["engine"]) != registered["version"]:
        raise ValueError("当前引擎版本未经本任务验收，请联系管理员核对冻结版本")
    with contextlib.redirect_stdout(sys.stderr):
        model = {
            "pandapower": prepare_flow,
            "pypsa": planning_network,
            "andes": prepare_dynamic,
        }[spec["engine"]](spec, directory)
    preserved = directory / ("model-input" + model.suffix)
    shutil.copyfile(model, preserved)
    spec["modelHash"] = hashlib.sha256(preserved.read_bytes()).hexdigest()
    transport = (
        spec.get("runSpec", {})
        .get("solver", {})
        .get("options", {})
        .get("transport", spec.get("transport", "mcp"))
    )
    if spec["engine"] == "pandapower" and transport == "direct":
        from analysis import static_analysis, failure

        event("solving", "标准 pandapower 后端计算中")
        try:
            with contextlib.redirect_stdout(sys.stderr):
                raw = static_analysis(str(model), spec["parameters"])
        except ValueError as e:
            raw = failure("solver_error", str(e))
    else:
        from mcp import ClientSession, StdioServerParameters
        from mcp.client.stdio import stdio_client

        env = os.environ.copy()
        env.update(
            {
                "POWERIO_MCP_ALLOWED_ROOTS": str(directory),
                "POWERMCP_HOME": str(directory / "powermcp"),
                "PYTHONUNBUFFERED": "1",
                "OMP_NUM_THREADS": "1",
                "OPENBLAS_NUM_THREADS": "1",
            }
        )
        (directory / "powermcp").mkdir(exist_ok=True)
        definition = json.loads((HERE / "servers.json").read_text())["servers"][
            spec["engine"]
        ]
        server = StdioServerParameters(
            command=sys.executable,
            args=[str(HERE / definition["entry"]), *definition["arguments"]],
            cwd=str(directory),
            env=env,
        )
        event("connecting", "连接独立 MCP 服务并发现工具")
        transcript = []
        with open(directory / "mcp-stderr.log", "w") as stderr:
            async with stdio_client(server, errlog=stderr) as (read, write):
                async with ClientSession(read, write) as session:
                    # These guards cover protocol startup only, never numerical solving.
                    init = await asyncio.wait_for(session.initialize(), timeout=180)
                    catalog = await asyncio.wait_for(session.list_tools(), timeout=120)
                    save(
                        directory / "tools.json",
                        {
                            "server": init.model_dump(mode="json"),
                            "tools": catalog.model_dump(mode="json"),
                        },
                    )
                    available = {tool.name for tool in catalog.tools}
                    if not set(definition["allowedTools"]) <= available:
                        raise ValueError(
                            "Required MCP tools are missing from the installed version"
                        )

                    async def call(name, args):
                        if name not in definition["allowedTools"]:
                            raise ValueError(
                                "Tool is not authorized for this calculation service"
                            )
                        transcript.append(
                            {"tool": name, "arguments": args, "startedAt": time.time()}
                        )
                        save(directory / "mcp.json", transcript)
                        reply = await session.call_tool(name, args)
                        value = unpack(reply)
                        transcript[-1].update(
                            {"finishedAt": time.time(), "result": value}
                        )
                        save(directory / "mcp.json", transcript)
                        return value

                    event("loading", "载入已校验的模型")
                    if spec["engine"] == "andes":
                        loaded = await call("run_power_flow", {"file_path": str(model)})
                        if not loaded.get("power_flow", {}).get(
                            "converged"
                        ) or not loaded.get("power_flow", {}).get(
                            "dynamic_models_loaded"
                        ):
                            raise ValueError(
                                "Initial power flow failed or model has no dynamics"
                            )
                    event("solving", "真实求解器计算中")
                    if spec["engine"] == "pandapower":
                        raw = await call(
                            "run_static_analysis",
                            {"file_path": str(model), "parameters": spec["parameters"]},
                        )
                    elif spec["engine"] == "pypsa":
                        raw = await call(
                            "run_planning_analysis",
                            {
                                "file_path": str(model),
                                "parameters": spec["parameters"],
                                "native_model": spec["modelId"] == "uploaded",
                            },
                        )
                    elif spec["parameters"].get("analysis") == "small_signal":
                        raw = await call("run_small_signal_analysis", {})
                    else:
                        raw = await call(
                            "run_time_domain_simulation",
                            {
                                "step_size": spec["parameters"]["stepS"],
                                "t_end": spec["parameters"]["endAtS"],
                            },
                        )
    event("checking", "核对求解状态、完整数据与业务判据")
    if raw.get("termination_condition") == "input_error":
        from analysis import failure

        raw = failure("solver_error", raw["message"])
    result = (
        raw
        if raw.get("solverStatus")
        or spec["engine"] == "pandapower"
        or (spec["engine"] == "pypsa" and spec["modelId"] == "uploaded")
        or (
            spec["engine"] == "andes"
            and spec["parameters"].get("analysis") == "small_signal"
        )
        else (planning_result if spec["engine"] == "pypsa" else dynamic_result)(
            spec, raw
        )
    )
    result.update(
        {
            "schemaVersion": "power-result/v1",
            "dataOrigin": "engine",
            "computedByEngine": True,
            "engine": spec["engine"],
            "modelId": spec["modelId"],
            "modelHash": spec["modelHash"],
            "inputHash": spec["inputHash"],
            "runId": spec["runId"],
            "versions": {
                p: importlib.metadata.version(p)
                for p in (
                    [spec["engine"]]
                    if transport == "direct"
                    else ["powermcp", "mcp", spec["engine"]]
                )
                + (["highspy"] if spec["engine"] == "pypsa" else [])
            },
        }
    )
    analysis = spec["parameters"].get(
        "analysis",
        "time_domain"
        if spec["engine"] == "andes"
        else "storage_dispatch"
        if spec["engine"] == "pypsa"
        else "ac_flow",
    )
    result.update(
        transport=transport,
        source="user"
        if spec.get("modelUploadId") or spec["modelId"] == "custom-timeseries"
        else "sample",
        version=result["versions"][spec["engine"]],
        solver="HiGHS"
        if spec["engine"] == "pypsa"
        else "ANDES EIG"
        if analysis == "small_signal"
        else "ANDES TDS"
        if spec["engine"] == "andes"
        else "PYPOWER OPF"
        if "opf" in analysis
        else "DC linear"
        if analysis == "dc_flow"
        else spec["parameters"].get("algorithm", "nr"),
        fidelity="linearized-dynamic"
        if analysis == "small_signal"
        else "nonlinear-dynamic"
        if spec["engine"] == "andes"
        else (
            "native-network-linear"
            if spec["modelId"] == "uploaded"
            else "single-bus-active-power"
        )
        if spec["engine"] == "pypsa"
        else "DC"
        if analysis.startswith("dc_")
        else "AC",
        execution=dict(
            startedAt=datetime.fromtimestamp(started, timezone.utc).isoformat(),
            finishedAt=datetime.now(timezone.utc).isoformat(),
            durationMs=round((time.time() - started) * 1000),
        ),
    )
    from result_tables import export_table

    counts = result.setdefault("tableCounts", {})
    for name, rows in result["tables"].items():
        if name not in counts:
            result["tables"][name], counts[name] = export_table(
                directory, name, iter(rows)
            )
    save(directory / "engine-result.json", result)
    save(directory / "input.json", spec.get("runSpec") or spec)
    event("saved", "结果、输入和原始调用记录已保存")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("directory")
    args = parser.parse_args()
    directory = Path(args.directory).resolve()
    spec = json.loads(Path(args.input).read_text())
    if spec.get("version") == "1.1" and "metadata" in spec:
        canonical = spec
        meta = canonical["metadata"]
        spec = dict(
            engine=canonical["solver"]["backend"],
            modelId=meta["model_id"],
            parameters=meta["parameters"],
            runId=meta["run_id"],
            inputHash=meta["input_digest"],
            runSpec=canonical,
            modelUploadId=meta.get("model_upload_id"),
        )
    if (
        spec["engine"] not in ["pandapower", "pypsa", "andes"]
        or Path(args.input).resolve().parent != directory
    ):
        raise ValueError("Invalid immutable run input")
    asyncio.run(execute(spec, directory))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        event("error", str(exc))
        raise
