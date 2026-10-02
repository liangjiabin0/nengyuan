#!/usr/bin/env python3
"""Build an SGCC guide proposal from the fixed DOCX template and structured JSON."""

from __future__ import annotations

import argparse
import copy
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ORDINALS = ["一是", "二是", "三是", "四是", "五是", "六是", "七是", "八是", "九是", "十是"]
REQUIRED_FIELDS = [
    "guide_year",
    "cover_date",
    "guide_group",
    "project_name",
    "discipline_root_technology",
    "plan_alignment",
    "application_scenario",
    "functional_requirements",
    "key_technologies",
    "topics",
    "maturity_and_application",
    "duration",
    "budget_limit",
    "lead_unit",
]


def fail(message: str) -> None:
    raise ValueError(message)


def clean_sentence(value: str) -> str:
    return value.strip().rstrip("；;。")


def ensure_period(value: str) -> str:
    value = value.strip()
    return value if value.endswith(("。", "！", "？")) else value + "。"


def numbered_items(items: list[str]) -> str:
    if not items or len(items) > len(ORDINALS):
        fail("Numbered lists must contain 1 to 10 items")
    cleaned = [clean_sentence(item) for item in items]
    if any(not item for item in cleaned):
        fail("Numbered list items must not be empty")
    return "；".join(f"{ORDINALS[index]}{item}" for index, item in enumerate(cleaned)) + "。"


def copy_run_properties(source_run, target_run) -> None:
    if source_run._r.rPr is not None:
        target_run._r.insert(0, copy.deepcopy(source_run._r.rPr))


def replace_text(paragraph, text: str) -> None:
    source_run = paragraph.runs[0] if paragraph.runs else None
    paragraph.clear()
    target_run = paragraph.add_run(text)
    if source_run is not None:
        copy_run_properties(source_run, target_run)


def replace_labelled(paragraph, label: str, text: str) -> None:
    source_runs = paragraph.runs
    label_source = source_runs[0] if source_runs else None
    body_source = source_runs[1] if len(source_runs) > 1 else label_source
    paragraph.clear()
    label_run = paragraph.add_run(label)
    body_run = paragraph.add_run(text)
    if label_source is not None:
        copy_run_properties(label_source, label_run)
    if body_source is not None:
        copy_run_properties(body_source, body_run)


def find_paragraph(document: Document, marker: str):
    matches = [paragraph for paragraph in document.paragraphs if marker in paragraph.text]
    if len(matches) != 1:
        fail(f"Template marker {marker!r} occurs {len(matches)} times")
    return matches[0]


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def paragraph_position(paragraphs, target) -> int:
    return next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is target._p)


def resize_topic_blocks(document: Document, topic_count: int) -> None:
    title_pattern = re.compile(r"^课题\d+：【课题\d+名称】$")
    title_paragraphs = [paragraph for paragraph in document.paragraphs if title_pattern.match(paragraph.text)]
    if len(title_paragraphs) != 3:
        fail("Fixed template must contain exactly three placeholder topic blocks")

    all_paragraphs = document.paragraphs
    last_start = paragraph_position(all_paragraphs, title_paragraphs[-1])
    source_block = [copy.deepcopy(paragraph._p) for paragraph in all_paragraphs[last_start:last_start + 5]]
    if len(source_block) != 5:
        fail("Fixed template topic block is incomplete")

    if topic_count < 3:
        for title_paragraph in reversed(title_paragraphs[topic_count:]):
            current = document.paragraphs
            start = paragraph_position(current, title_paragraph)
            for paragraph in reversed(current[start:start + 5]):
                remove_paragraph(paragraph)
    elif topic_count > 3:
        tail = find_paragraph(document, "四、技术成熟度及成果应用")
        for _ in range(topic_count - 3):
            for element in source_block:
                tail._p.addprevious(copy.deepcopy(element))


def validate_payload(data: dict[str, Any]) -> None:
    missing = [field for field in REQUIRED_FIELDS if field not in data]
    if missing:
        fail(f"Missing required fields: {missing}")

    functional = data["functional_requirements"]
    technologies = data["key_technologies"]
    topics = data["topics"]
    if not isinstance(functional, list) or not isinstance(technologies, list) or not isinstance(topics, list):
        fail("functional_requirements, key_technologies and topics must be arrays")
    if not topics:
        fail("At least one topic is required")
    if not (len(functional) == len(technologies) == len(topics)):
        fail("Functional requirement count, key technology count and topic count must be equal")

    for index, topic in enumerate(topics, start=1):
        for field in ("title", "problem", "research_points", "metrics"):
            if field not in topic:
                fail(f"Topic {index} is missing {field}")
        points = topic["research_points"]
        if not isinstance(points, list) or len(points) != 3:
            fail(f"Topic {index} must contain exactly three research_points")
        metrics = topic["metrics"]
        if not isinstance(metrics, dict):
            fail(f"Topic {index} metrics must be an object")
        for metric in ("performance", "validation", "outputs"):
            if not str(metrics.get(metric, "")).strip():
                fail(f"Topic {index} is missing metrics.{metric}")

    serialized = json.dumps(data, ensure_ascii=False)
    if "【" in serialized or "】" in serialized:
        fail("Payload must not contain template placeholder brackets")


def request_field_update(document: Document) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def build(data: dict[str, Any], template: Path, output: Path) -> None:
    validate_payload(data)
    document = Document(template)
    topics = data["topics"]
    resize_topic_blocks(document, len(topics))

    guide_year = str(data["guide_year"]).strip()
    project_number = int(data.get("project_number", 1))
    project_page = int(data.get("project_page", 2))
    guide_group = str(data["guide_group"]).strip()
    project_name = str(data["project_name"]).strip()

    replace_text(find_paragraph(document, "【指南年度】"), f"国家电网公司{guide_year}年")
    replace_text(find_paragraph(document, "【编制年月】"), str(data["cover_date"]).strip())
    replace_text(find_paragraph(document, "【目录指南组】"), f"{guide_group}（1项指南）………………1")
    replace_text(find_paragraph(document, "【目录项目名称】"), f"{project_name}…………{project_page}")
    replace_text(find_paragraph(document, "【正文指南组】"), f"{guide_group}（1项指南）")
    replace_text(find_paragraph(document, "【正文项目名称】"), f"{project_number}. {project_name}")
    replace_text(find_paragraph(document, "【学科与根技术】"), str(data["discipline_root_technology"]).strip())
    replace_text(find_paragraph(document, "【规划对应关系】"), str(data["plan_alignment"]).strip())
    replace_labelled(find_paragraph(document, "【应用场景】"), "应用场景：", str(data["application_scenario"]).strip())
    replace_labelled(
        find_paragraph(document, "【功能需求】"),
        "功能需求：",
        numbered_items([str(item) for item in data["functional_requirements"]]),
    )
    key_text = numbered_items([str(item) for item in data["key_technologies"]])
    route = str(data.get("technical_route", "")).strip()
    if route:
        key_text += ensure_period(route)
    replace_labelled(find_paragraph(document, "【关键技术】"), "关键技术：", key_text)

    title_pattern = re.compile(r"^课题\d+：【课题\d+名称】$")
    title_paragraphs = [paragraph for paragraph in document.paragraphs if title_pattern.match(paragraph.text)]
    if len(title_paragraphs) != len(topics):
        fail("Topic block resizing failed")
    for index, (title_paragraph, topic) in enumerate(zip(title_paragraphs, topics), start=1):
        current = document.paragraphs
        start = paragraph_position(current, title_paragraph)
        block = current[start:start + 5]
        if len(block) != 5:
            fail(f"Topic {index} block is incomplete")
        replace_text(block[0], f"课题{index}：{str(topic['title']).strip()}")
        goal = f"针对{clean_sentence(str(topic['problem']))}，{numbered_items([str(point) for point in topic['research_points']])}"
        interface = str(topic.get("interface", "")).strip()
        if interface:
            goal += ensure_period(interface)
        replace_labelled(block[1], "预期目标：", goal)
        metrics = topic["metrics"]
        replace_labelled(block[2], "考核指标：", f"（1）{ensure_period(str(metrics['performance']))}")
        replace_text(block[3], f"（2）{ensure_period(str(metrics['validation']))}")
        replace_text(block[4], f"（3）{ensure_period(str(metrics['outputs']))}")

    replace_text(find_paragraph(document, "【成熟度与成果应用】"), str(data["maturity_and_application"]).strip())
    replace_text(find_paragraph(document, "【项目实施期限】"), str(data["duration"]).strip())
    replace_text(find_paragraph(document, "【支持经费限额】"), str(data["budget_limit"]).strip())
    replace_text(find_paragraph(document, "【牵头单位】"), str(data["lead_unit"]).strip())

    document.core_properties.title = project_name
    document.core_properties.subject = f"国家电网公司{guide_year}年总部科技项目申报指南（建议稿）"
    request_field_update(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path, help="UTF-8 JSON payload")
    parser.add_argument("--output", required=True, type=Path, help="Output DOCX")
    parser.add_argument(
        "--template",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "assets" / "sgcc-guide-fixed-template.docx",
    )
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8-sig"))
    build(data, args.template, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
