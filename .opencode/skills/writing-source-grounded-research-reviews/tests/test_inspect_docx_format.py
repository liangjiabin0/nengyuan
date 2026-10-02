import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from scripts.inspect_docx_format import inspect_docx


def set_run_fonts(run, east_asia, latin):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:cs"), latin)


class DocxFormatInspectorTests(unittest.TestCase):
    def test_detects_bad_heading_fonts_and_grid_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.docx"
            doc = Document()
            h = doc.add_heading("1 测试Heading", level=1)
            h.runs[0].font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
            set_run_fonts(h.runs[0], "等线", "Calibri")
            p = doc.add_paragraph()
            run = p.add_run("中文ABC123")
            set_run_fonts(run, "微软雅黑", "Calibri")
            table = doc.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            table.cell(0, 0).text = "表头"
            table.cell(0, 1).text = "Header"
            doc.save(path)

            report = inspect_docx(path)
            self.assertGreater(report["summary"]["heading_color_issues"], 0)
            self.assertGreater(report["summary"]["font_issues"], 0)
            self.assertGreater(report["summary"]["table_issues"], 0)
            self.assertEqual(report["summary"]["native_equation_count"], 0)

    def test_accepts_basic_font_split_and_superscript(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "good.docx"
            doc = Document()
            p = doc.add_paragraph()
            r1 = p.add_run("中文")
            set_run_fonts(r1, "宋体", "Times New Roman")
            r2 = p.add_run("ABC123")
            set_run_fonts(r2, "宋体", "Times New Roman")
            r3 = p.add_run("[1]")
            set_run_fonts(r3, "宋体", "Times New Roman")
            r3.font.superscript = True
            doc.save(path)

            report = inspect_docx(path)
            self.assertEqual(report["summary"]["font_issues"], 0)
            self.assertEqual(report["summary"]["citation_superscript_issues"], 0)


if __name__ == "__main__":
    unittest.main()
