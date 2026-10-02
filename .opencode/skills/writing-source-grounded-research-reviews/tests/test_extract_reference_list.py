import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.extract_reference_list import extract_references, parse_reference_entry, read_document_text


class ReferenceExtractionTests(unittest.TestCase):
    def test_parse_chinese_journal_reference(self):
        entry = "[2] 罗佳怡, 邢海军, 全文斌, 施怡沁. 极端高温下新型电力系统韧性评估及差异化规划[J]. 电力系统自动化, 2025, 49(3):135-144."
        record = parse_reference_entry(entry)
        self.assertEqual(record["index"], 2)
        self.assertEqual(record["year"], 2025)
        self.assertEqual(record["authors"][0], "罗佳怡")
        self.assertEqual(record["title"], "极端高温下新型电力系统韧性评估及差异化规划")
        self.assertIn("电力系统自动化", record["source"])

    def test_parse_online_first_reference(self):
        entry = "[1] 江波, 雷洪涛, 刘亚杰. 互联微电网韧性评估、调度与重构优化研究综述[J/OL]. 电工技术学报, 2026-01-04网络首发. DOI:10.19595/example."
        record = parse_reference_entry(entry)
        self.assertEqual(record["year"], 2026)
        self.assertEqual(record["document_type"], "J/OL")
        self.assertEqual(record["doi"], "10.19595/example")

    def test_malformed_entry_is_conservative(self):
        record = parse_reference_entry("[7] 作者信息缺失且没有年份")
        self.assertEqual(record["index"], 7)
        self.assertIsNone(record["year"])
        self.assertIn("year_not_found", record["warnings"])

    def test_extract_reference_section_from_text(self):
        text = "正文引用[1]。\n参考文献\n[1] 张三. 测试题名[J]. 测试期刊, 2024, 1(1):1-2.\n[2] 李四. 第二题名[D]. 北京: 某大学, 2023.\n附录"
        records = extract_references(text)
        self.assertEqual([r["index"] for r in records], [1, 2])

    def test_uses_last_reference_heading_after_table_of_contents(self):
        text = (
            "目录\n1 正文\n参考文献\n2. 正文章节\n正文结论[1]。\n"
            "参考文献\n[1] 张三. 真实题名[J]. 期刊, 2024.\n"
        )
        records = extract_references(text)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["title"], "真实题名")

    def test_extracts_from_directory_and_records_source_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "a.txt").write_text(
                "参考文献\n[1] 张三. 题名A[J]. 期刊, 2024.", encoding="utf-8"
            )
            (root / "b.md").write_text(
                "参考文献\n[2] 李四. 题名B[J]. 期刊, 2023.", encoding="utf-8"
            )
            records = extract_references(root)
            self.assertEqual(len(records), 2)
            self.assertEqual({r["source_file"] for r in records}, {"a.txt", "b.md"})

    def test_read_docx_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"
            doc = Document()
            doc.add_paragraph("参考文献")
            doc.add_paragraph("[1] 张三. 测试题名[J]. 测试期刊, 2024, 1(1):1-2.")
            doc.save(path)
            text = read_document_text(path)
            self.assertIn("测试题名", text)


if __name__ == "__main__":
    unittest.main()
