#!/usr/bin/env python3
"""Extract and conservatively parse reference-list entries from DOCX/text files.

The parser is intentionally conservative. It produces warnings instead of
inventing missing bibliographic fields. Always verify records against the
source document's title page or reference list before publication.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Iterable, Sequence

try:
    from docx import Document
except ImportError:  # pragma: no cover - handled by CLI error
    Document = None  # type: ignore[assignment]

ENTRY_START_RE = re.compile(r"^\s*(?:\[(?P<bracket>\d+)\]|(?P<plain>\d+)[\.、])\s*(?P<body>.*)$")
YEAR_RE = re.compile(r"(?<!\d)((?:19|20)\d{2})(?!\d)")
DOCTYPE_RE = re.compile(r"\[([A-Za-z]+(?:/[A-Za-z]+)?)\]")
DOI_RE = re.compile(r"\bdoi\s*:\s*([^\s,;，；]+)", re.IGNORECASE)
REFERENCE_HEADING_RE = re.compile(r"^\s*(参考文献|References|Bibliography)\s*$", re.IGNORECASE)
STOP_HEADING_RE = re.compile(r"^\s*(附录|致谢|作者简介|Appendix|Acknowledg(?:e)?ments?)\b", re.IGNORECASE)


def read_document_text(path: str | Path) -> str:
    """Read DOCX, TXT, MD, HTML, or legacy text-like files as normalized text."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        if Document is None:
            raise RuntimeError("python-docx is required to read .docx files")
        doc = Document(str(file_path))
        lines: list[str] = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if suffix in {".txt", ".md", ".html", ".htm", ".doc"}:
        data = file_path.read_bytes()
        for encoding in ("utf-8", "utf-8-sig", "gb18030", "big5"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported input type: {suffix or '<no extension>'}")


def _split_author_and_rest(body: str) -> tuple[str, str]:
    match = re.search(r"[。.]\s*", body)
    if not match:
        return body.strip(), ""
    return body[: match.start()].strip(), body[match.end() :].strip()


def _clean_doi(raw: str | None) -> str | None:
    if not raw:
        return None
    return raw.rstrip(".。;,；")


def parse_reference_entry(entry: str) -> dict:
    """Parse a single sequential reference entry into a JSON-safe record."""
    raw = " ".join(entry.strip().split())
    warnings: list[str] = []
    start = ENTRY_START_RE.match(raw)
    index: int | None = None
    body = raw
    if start:
        index = int(start.group("bracket") or start.group("plain"))
        body = start.group("body").strip()
    else:
        warnings.append("reference_number_not_found")

    author_text, rest = _split_author_and_rest(body)
    authors = [a.strip() for a in re.split(r"[,，;；]", author_text) if a.strip()]
    if not authors:
        warnings.append("authors_not_found")

    doctype_match = DOCTYPE_RE.search(rest)
    document_type = doctype_match.group(1).upper() if doctype_match else None
    if doctype_match:
        title = rest[: doctype_match.start()].strip(" .。")
        source = rest[doctype_match.end() :].strip(" .。")
    else:
        # Fallback: use the next sentence as title when no type marker exists.
        title_part, source = _split_author_and_rest(rest)
        title = title_part.strip(" .。")
        if not title:
            title = rest.strip(" .。")
            source = ""
        warnings.append("document_type_not_found")

    year_matches = YEAR_RE.findall(source or rest)
    year = int(year_matches[0]) if year_matches else None
    if year is None:
        warnings.append("year_not_found")

    doi_match = DOI_RE.search(source or rest)
    doi = _clean_doi(doi_match.group(1) if doi_match else None)

    if not title:
        warnings.append("title_not_found")
    if not source:
        warnings.append("source_not_found")

    return {
        "index": index,
        "authors": authors,
        "author_text": author_text or None,
        "year": year,
        "title": title or None,
        "document_type": document_type,
        "source": source or None,
        "doi": doi,
        "raw": raw,
        "warnings": warnings,
    }


def _reference_lines(text: str) -> list[str]:
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    heading_indices = [i for i, line in enumerate(lines) if REFERENCE_HEADING_RE.match(line)]
    heading_index = heading_indices[-1] if heading_indices else None
    scan = lines[heading_index + 1 :] if heading_index is not None else lines

    entries: list[str] = []
    current: list[str] = []
    for line in scan:
        if not line:
            continue
        if STOP_HEADING_RE.match(line) and current:
            break
        if ENTRY_START_RE.match(line):
            if current:
                entries.append(" ".join(current))
            current = [line]
        elif current:
            # Wrapped bibliographic lines are joined until the next numbered entry.
            current.append(line)
    if current:
        entries.append(" ".join(current))
    return entries


def _existing_path(value: str | Path) -> Path | None:
    if isinstance(value, Path):
        return value if value.exists() else None
    try:
        candidate = Path(value)
        return candidate if candidate.exists() else None
    except (OSError, ValueError):
        return None


def extract_references(source: str | Path) -> list[dict]:
    """Extract parsed references from raw text, a file, or a directory."""
    path = _existing_path(source)
    if path is not None and path.is_dir():
        supported = {".docx", ".txt", ".md", ".html", ".htm", ".doc"}
        records: list[dict] = []
        for child in sorted(p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in supported):
            for record in extract_references(child):
                enriched = dict(record)
                enriched["source_file"] = str(child.relative_to(path))
                records.append(enriched)
        return records
    if path is not None:
        text = read_document_text(path)
    else:
        text = str(source)
    return [parse_reference_entry(entry) for entry in _reference_lines(text)]


def _write_json(records: Sequence[dict], path: Path) -> None:
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_csv(records: Sequence[dict], path: Path) -> None:
    fieldnames = [
        "index",
        "authors",
        "year",
        "title",
        "document_type",
        "source",
        "doi",
        "warnings",
        "raw",
        "source_file",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            row = dict(record)
            row["authors"] = "; ".join(record.get("authors", []))
            row["warnings"] = "; ".join(record.get("warnings", []))
            writer.writerow({key: row.get(key) for key in fieldnames})


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="DOCX/TXT/MD/HTML file or directory containing reference lists")
    parser.add_argument("--json", dest="json_path", type=Path, help="Write records as UTF-8 JSON")
    parser.add_argument("--csv", dest="csv_path", type=Path, help="Write records as UTF-8 BOM CSV")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    records = extract_references(args.input)
    if args.json_path:
        _write_json(records, args.json_path)
    if args.csv_path:
        _write_csv(records, args.csv_path)
    if not args.json_path and not args.csv_path:
        print(json.dumps(records, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
