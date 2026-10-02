#!/usr/bin/env python3
"""Render canonical paper-review JSON to DOCX, optionally using a review DOCX as a style/layout template."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ZH = {
    "report": "论文审稿意见",
    "overall": "一、总体评价",
    "experiment": "二、核心结论—实验证据矩阵",
    "repro": "三、复现性审查",
    "latest": "四、最新相关文献与对比审查",
    "theory": "五、理论分析审查",
    "major": "六、主要审稿意见",
    "minor": "七、次要意见",
    "decision": "八、审稿结论",
    "claim": "核心结论", "evidence": "现有证据", "status": "支撑状态", "judgment": "判断",
    "item": "复现要素", "loc": "位置", "req": "审稿要求",
    "work": "近期代表性工作", "relation": "相关性/对比情况", "year": "年份", "source": "来源",
    "theory_claim": "理论结论", "derivation": "推导/证明", "assumptions": "关键假设", "assessment": "判断",
    "recommendation": "建议",
}
EN = {
    "report": "Peer Review Comments",
    "overall": "1. Overall Assessment",
    "experiment": "2. Central Claim–Experimental Evidence Matrix",
    "repro": "3. Reproducibility Audit",
    "latest": "4. Latest Relevant Literature and Comparator Audit",
    "theory": "5. Theoretical Analysis Audit",
    "major": "6. Major Comments",
    "minor": "7. Minor Comments",
    "decision": "8. Recommendation",
    "claim": "Central claim", "evidence": "Evidence", "status": "Support status", "judgment": "Assessment",
    "item": "Reproducibility item", "loc": "Location", "req": "Review requirement",
    "work": "Recent relevant work", "relation": "Relevance / comparison", "year": "Year", "source": "Source",
    "theory_claim": "Theoretical claim", "derivation": "Derivation / proof", "assumptions": "Key assumptions", "assessment": "Assessment",
    "recommendation": "Recommendation",
}



STATUS_ZH = {
    "supported": "充分支撑", "partially_supported": "部分支撑", "unsupported": "不支撑", "not_assessable": "无法判断",
    "ready": "可复现", "minor_gaps": "存在少量缺口", "major_gaps": "存在关键缺口", "not_applicable": "不适用",
    "reporting_only": "仅审查报告完整性", "artifacts_inspected": "已检查工件", "partial_execution": "已部分执行",
    "end_to_end_execution": "已端到端执行", "adequate": "充分", "partially_adequate": "部分充分", "inadequate": "不足",
    "not_required": "无需理论证明", "required": "需要理论支撑", "optional": "可选",
    "specified": "已说明", "partially_specified": "部分说明", "missing": "缺失", "inaccessible": "不可访问",
    "accept": "接收", "minor_revision": "小修", "major_revision": "大修", "reject": "拒稿", "unable_to_assess": "无法评估"
}
STATUS_EN = {
    "supported": "Supported", "partially_supported": "Partially supported", "unsupported": "Unsupported", "not_assessable": "Not assessable",
    "ready": "Ready", "minor_gaps": "Minor gaps", "major_gaps": "Major gaps", "not_applicable": "Not applicable",
    "reporting_only": "Reporting only", "artifacts_inspected": "Artifacts inspected", "partial_execution": "Partial execution",
    "end_to_end_execution": "End-to-end execution", "adequate": "Adequate", "partially_adequate": "Partially adequate", "inadequate": "Inadequate",
    "not_required": "Not required", "required": "Required", "optional": "Optional",
    "specified": "Specified", "partially_specified": "Partially specified", "missing": "Missing", "inaccessible": "Inaccessible",
    "accept": "Accept", "minor_revision": "Minor revision", "major_revision": "Major revision", "reject": "Reject", "unable_to_assess": "Unable to assess"
}

def status_label(value, lang: str) -> str:
    raw = str(value or "")
    return (STATUS_ZH if lang == "zh-CN" else STATUS_EN).get(raw, raw.replace("_", " "))

def refresh_template_headers_footers(doc: Document, lang: str) -> None:
    tagline = (
        "论文审稿意见 | 四项核心审查：实验支撑、可复现性、最新工作对比、理论分析"
        if lang == "zh-CN" else
        "Peer Review Comments | Core audits: experimental support, reproducibility, latest comparators, theoretical analysis"
    )
    for section in doc.sections:
        # Remove stale manuscript-specific textual header content while preserving field-only paragraphs.
        for p in section.header.paragraphs:
            if p.text.strip():
                p.clear()
        written = False
        for p in section.footer.paragraphs:
            if p.text.strip():
                p.clear()
                if not written:
                    p.add_run(tagline)
                    written = True
        if not written and section.footer.paragraphs:
            section.footer.paragraphs[0].add_run(tagline)

def clear_body_keep_sections(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_default_fonts(doc: Document, lang: str) -> None:
    if "Normal" not in doc.styles:
        return
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    if lang == "zh-CN":
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    else:
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    try:
        p = doc.add_heading(text, level=level)
    except Exception:
        p = doc.add_paragraph(text)
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(14 if level == 1 else 12)
    p.paragraph_format.keep_with_next = True


def add_title(doc: Document, report_title: str, paper_title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(report_title)
    r.bold = True
    r.font.size = Pt(16)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"《{paper_title}》" if report_title.startswith("论文") else paper_title)
    r2.bold = True
    r2.font.size = Pt(14)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = h
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value or ""
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "; ".join(str(x) for x in value)
    return str(value)


def render(data: dict, output: Path, template: Path | None = None) -> None:
    lang = data.get("metadata", {}).get("review_language", "en")
    ui = ZH if lang == "zh-CN" else EN
    doc = Document(str(template)) if template else Document()
    if template:
        clear_body_keep_sections(doc)
        refresh_template_headers_footers(doc, lang)
    set_default_fonts(doc, lang)

    paper_title = data.get("metadata", {}).get("paper_title", "Untitled")
    add_title(doc, ui["report"], paper_title)

    add_heading(doc, ui["overall"])
    doc.add_paragraph(data.get("summary", ""))
    rationale = data.get("verdict", {}).get("rationale", "")
    if rationale:
        doc.add_paragraph(rationale)

    add_heading(doc, ui["experiment"])
    cs = data.get("conclusion_support", {})
    rows = []
    for c in cs.get("claims", []):
        rows.append([
            f"{c.get('id','')} {c.get('claim','')}",
            to_text(c.get("supporting_evidence")),
            status_label(c.get("support_status", ""), lang),
            c.get("evidence_assessment", "") or c.get("missing_or_weak_evidence", ""),
        ])
    add_table(doc, [ui["claim"], ui["evidence"], ui["status"], ui["judgment"]], rows)

    add_heading(doc, ui["repro"])
    repro = data.get("reproducibility", {})
    if repro:
        doc.add_paragraph(f"{status_label(repro.get('reproduction_readiness',''), lang)}: {repro.get('rationale','')}")
    rows = []
    for x in repro.get("checked_items", []):
        rows.append([x.get("item", ""), status_label(x.get("status", ""), lang), x.get("location", ""), x.get("evidence", "")])
    add_table(doc, [ui["item"], ui["status"], ui["loc"], ui["req"]], rows)

    add_heading(doc, ui["latest"])
    latest = data.get("latest_literature_comparison", {})
    if latest:
        doc.add_paragraph(f"{status_label(latest.get('comparison_status',''), lang)}: {latest.get('rationale','')}")
    rows = []
    for w in latest.get("recent_relevant_works", []):
        relation = w.get("relevance", "")
        relation += f"; cited={w.get('cited_in_manuscript')}; compared={w.get('experimentally_compared')}"
        rows.append([w.get("title", ""), to_text(w.get("year")), w.get("source", ""), relation])
    add_table(doc, [ui["work"], ui["year"], ui["source"], ui["relation"]], rows)
    if latest.get("required_resolution"):
        doc.add_paragraph(latest["required_resolution"])

    add_heading(doc, ui["theory"])
    theory = data.get("theoretical_analysis", {})
    if theory:
        doc.add_paragraph(f"{status_label(theory.get('overall_status',''), lang)}: {theory.get('rationale','')}")
    rows = []
    for t in theory.get("claims", []):
        rows.append([
            f"{t.get('id','')} {t.get('claim','')}",
            t.get("supporting_derivation", ""),
            to_text(t.get("assumptions")),
            f"{status_label(t.get('theoretical_support_status',''), lang)}: {t.get('assessment','')}",
        ])
    add_table(doc, [ui["theory_claim"], ui["derivation"], ui["assumptions"], ui["assessment"]], rows)

    issues = data.get("issues", [])
    major = [x for x in issues if x.get("severity") in {"critical", "major", "moderate"}]
    minor = [x for x in issues if x.get("severity") == "minor"]

    add_heading(doc, ui["major"])
    if major:
        for i, issue in enumerate(major, 1):
            p = doc.add_paragraph(style="List Number" if "List Number" in doc.styles else None)
            r = p.add_run(issue.get("title", ""))
            r.bold = True
            doc.add_paragraph(issue.get("problem", ""))
            if issue.get("required_action"):
                doc.add_paragraph(issue["required_action"])
    else:
        doc.add_paragraph("无" if lang == "zh-CN" else "None")

    add_heading(doc, ui["minor"])
    if minor:
        for issue in minor:
            doc.add_paragraph(issue.get("title", "") + (": " + issue.get("problem", "") if issue.get("problem") else ""), style="List Bullet" if "List Bullet" in doc.styles else None)
    else:
        doc.add_paragraph("无" if lang == "zh-CN" else "None")

    add_heading(doc, ui["decision"])
    decision = data.get("verdict", {}).get("decision", "")
    doc.add_paragraph(f"{ui['recommendation']}: {status_label(decision, lang)}")
    if rationale:
        doc.add_paragraph(rationale)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--output", required=True, type=Path)
    ap.add_argument("--template", type=Path)
    args = ap.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8-sig"))
    render(data, args.output, args.template)
    print(f"Rendered DOCX: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
