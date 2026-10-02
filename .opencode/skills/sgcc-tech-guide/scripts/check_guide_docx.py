#!/usr/bin/env python3
"""Run static OOXML checks on an SGCC guide DOCX before visual review."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SECTIONS = [
    "一、学科、根技术",
    "二、应用场景、功能需求及关键技术",
    "三、课题内容",
    "四、技术成熟度及成果应用",
    "五、项目实施期限",
    "六、支持经费限额",
    "七、牵头单位",
]
LABELS = ["应用场景：", "功能需求：", "关键技术：", "预期目标：", "考核指标："]
BODY_FONT = "仿宋_GB2312"
LATIN_FONT = "Times New Roman"


def run_east_asia_font(run) -> str | None:
    if run._r.rPr is None or run._r.rPr.rFonts is None:
        return None
    return run._r.rPr.rFonts.get(qn("w:eastAsia"))


def run_latin_fonts(run) -> tuple[str | None, str | None, str | None]:
    if run._r.rPr is None or run._r.rPr.rFonts is None:
        return None, None, None
    fonts = run._r.rPr.rFonts
    return fonts.get(qn("w:ascii")), fonts.get(qn("w:hAnsi")), fonts.get(qn("w:cs"))


def first_nonempty_run(paragraph):
    return next((run for run in paragraph.runs if run.text), None)


def close(actual: float | None, expected: float, tolerance: float = 0.03) -> bool:
    return actual is not None and abs(actual - expected) <= tolerance


def paragraph_index(paragraphs, text: str) -> int:
    return next((index for index, paragraph in enumerate(paragraphs) if paragraph.text == text), -1)


def check_fixed_format(document: Document) -> tuple[list[str], dict[str, object]]:
    errors: list[str] = []
    paragraphs = document.paragraphs
    texts = [paragraph.text for paragraph in paragraphs]
    details: dict[str, object] = {}

    if len(document.sections) != 3:
        errors.append(f"Fixed format requires 3 sections, found {len(document.sections)}")
    if document.tables:
        errors.append(f"Fixed format forbids Word tables, found {len(document.tables)}")

    expected_section_cm = {
        "page_width": 21.0,
        "page_height": 29.7,
        "top_margin": 2.80,
        "bottom_margin": 2.50,
        "left_margin": 2.80,
        "right_margin": 2.60,
        "header_distance": 1.50,
        "footer_distance": 1.60,
    }
    section_mismatches: list[str] = []
    for index, section in enumerate(document.sections, start=1):
        for field, expected in expected_section_cm.items():
            value = getattr(section, field)
            actual = value.cm if value is not None else None
            if not close(actual, expected):
                section_mismatches.append(f"section {index} {field}={actual}")
    if section_mismatches:
        errors.append("Page setup differs from fixed template: " + "; ".join(section_mismatches))
    details["page_setup_mismatches"] = section_mismatches

    fixed_text_checks = {
        0: "附件2",
        7: "总部科技项目申报指南（建议稿）",
        17: "目    录",
        30: "一、学科、根技术",
        33: "二、应用场景、功能需求及关键技术",
        37: "三、课题内容",
    }
    for index, expected in fixed_text_checks.items():
        if index >= len(texts) or texts[index] != expected:
            errors.append(f"Fixed paragraph {index} must be {expected!r}")
    if len(texts) <= 6 or not re.fullmatch(r"国家电网公司\d{4}年", texts[6]):
        errors.append("Cover year line does not match fixed format")
    if len(texts) <= 15 or not re.fullmatch(r"\d{4}年\d{1,2}月", texts[15]):
        errors.append("Cover date must use YYYY年M月")

    section_positions = [paragraph_index(paragraphs, name) for name in SECTIONS]
    if any(index < 0 for index in section_positions) or section_positions != sorted(section_positions):
        errors.append("Seven first-level sections are missing or out of order")

    topic_pattern = re.compile(r"^课题(\d+)：.+")
    topic_starts = [index for index, text in enumerate(texts) if topic_pattern.fullmatch(text)]
    if not topic_starts:
        errors.append("Fixed format requires at least one topic block")
    topic_errors: list[str] = []
    for expected_number, start in enumerate(topic_starts, start=1):
        match = topic_pattern.fullmatch(texts[start])
        actual_number = int(match.group(1)) if match else -1
        if actual_number != expected_number:
            topic_errors.append(f"topic number {actual_number} at block {expected_number}")
        if start + 4 >= len(texts):
            topic_errors.append(f"topic {expected_number} block is incomplete")
            continue
        expected_prefixes = ["预期目标：", "考核指标：（1）", "（2）", "（3）"]
        for offset, prefix in enumerate(expected_prefixes, start=1):
            if not texts[start + offset].startswith(prefix):
                topic_errors.append(f"topic {expected_number} paragraph {offset + 1} must start with {prefix}")
        goal = texts[start + 1]
        ordinal_positions = [goal.find(marker) for marker in ("一是", "二是", "三是")]
        if any(position < 0 for position in ordinal_positions) or ordinal_positions != sorted(ordinal_positions):
            topic_errors.append(f"topic {expected_number} goal must contain ordered 一是/二是/三是")
    if topic_starts:
        after_last = topic_starts[-1] + 5
        if after_last >= len(texts) or texts[after_last] != "四、技术成熟度及成果应用":
            topic_errors.append("Only five-paragraph topic blocks may appear before section four")
    if topic_errors:
        errors.append("Topic block structure differs from fixed format: " + "; ".join(topic_errors))
    details["topic_count"] = len(topic_starts)
    details["topic_errors"] = topic_errors

    for required_label in ("应用场景：", "功能需求：", "关键技术："):
        count = sum(text.startswith(required_label) for text in texts)
        if count != 1:
            errors.append(f"{required_label} must occur exactly once, found {count}")

    labelled_paragraphs = [paragraph for paragraph in paragraphs if any(paragraph.text.startswith(label) for label in LABELS)]
    label_errors: list[str] = []
    for paragraph in labelled_paragraphs:
        label = next(label for label in LABELS if paragraph.text.startswith(label))
        if len(paragraph.runs) < 2 or paragraph.runs[0].text != label:
            label_errors.append(f"{label} is not isolated in its own run")
            continue
        if paragraph.runs[0].bold is not True:
            label_errors.append(f"{label} is not bold")
        if paragraph.runs[1].bold is True:
            label_errors.append(f"body after {label} must not be bold")
    if label_errors:
        errors.append("Label formatting differs from fixed template: " + "; ".join(label_errors))
    details["label_errors"] = label_errors

    style_expectations = [
        (0, "黑体", 16.0, False),
        (6, "华文中宋", 22.0, False),
        (7, "华文中宋", 22.0, False),
        (15, BODY_FONT, 16.0, False),
        (17, "黑体", 20.0, False),
        (18, "黑体", 15.0, False),
        (19, "宋体", 14.0, False),
    ]
    style_errors: list[str] = []
    for index, expected_font, expected_size, expected_bold in style_expectations:
        if index >= len(paragraphs):
            continue
        run = first_nonempty_run(paragraphs[index])
        if run is None:
            style_errors.append(f"paragraph {index} has no formatted run")
            continue
        size = run.font.size.pt if run.font.size else None
        if run_east_asia_font(run) != expected_font or not close(size, expected_size) or bool(run.bold) != expected_bold:
            style_errors.append(
                f"paragraph {index}: font={run_east_asia_font(run)}, size={size}, bold={run.bold}"
            )

    group_candidates = [paragraph for paragraph in paragraphs if re.fullmatch(r".+组（1项指南）", paragraph.text)]
    if len(group_candidates) != 1:
        style_errors.append("body guide-group heading must occur exactly once")
    else:
        run = first_nonempty_run(group_candidates[0])
        size = run.font.size.pt if run and run.font.size else None
        if run is None or run_east_asia_font(run) != "黑体" or not close(size, 22.0):
            style_errors.append("body guide-group heading must be 22 pt Heiti")

    project_titles = [paragraph for paragraph in paragraphs if re.fullmatch(r"\d+\. .+", paragraph.text)]
    if len(project_titles) != 1:
        style_errors.append("numbered body project title must occur exactly once")
    else:
        run = first_nonempty_run(project_titles[0])
        size = run.font.size.pt if run and run.font.size else None
        if run is None or run_east_asia_font(run) != "黑体" or not close(size, 18.0) or run.bold is not True:
            style_errors.append("body project title must be bold 18 pt Heiti")

    heading_texts = set(SECTIONS)
    task_title_indices = set(topic_starts)
    for index, paragraph in enumerate(paragraphs):
        if not paragraph.text or index < 30:
            continue
        run = first_nonempty_run(paragraph)
        if run is None:
            continue
        size = run.font.size.pt if run.font.size else None
        if paragraph.text in heading_texts or index in task_title_indices:
            if run_east_asia_font(run) != "黑体" or not close(size, 16.0):
                style_errors.append(f"heading paragraph {index} must be 16 pt Heiti")
        else:
            if run_east_asia_font(run) != BODY_FONT or not close(size, 16.0):
                style_errors.append(f"body paragraph {index} must be 16 pt FangSong_GB2312")
            line_spacing = paragraph.paragraph_format.line_spacing
            line_pt = line_spacing.pt if hasattr(line_spacing, "pt") else None
            first_indent = paragraph.paragraph_format.first_line_indent
            indent_pt = first_indent.pt if first_indent is not None else None
            if not close(line_pt, 28.0) or not close(indent_pt, 32.0):
                style_errors.append(
                    f"body paragraph {index} must use 28 pt line spacing and 32 pt first-line indent"
                )
        for body_run in (item for item in paragraph.runs if item.text):
            ascii_font, hansi_font, cs_font = run_latin_fonts(body_run)
            if (ascii_font, hansi_font, cs_font) != (LATIN_FONT, LATIN_FONT, LATIN_FONT):
                style_errors.append(f"paragraph {index} Latin font slots must be Times New Roman")
                break
    if style_errors:
        errors.append("Typography differs from fixed template: " + "; ".join(style_errors))
    details["style_errors"] = style_errors

    placeholder_count = sum(text.count("【") + text.count("】") for text in texts)
    if placeholder_count:
        errors.append(f"Template placeholders remain: {placeholder_count}")
    details["placeholder_bracket_count"] = placeholder_count
    return errors, details


def read_zip_parts(path: Path) -> tuple[dict[str, bytes], str | None]:
    try:
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            if bad:
                return {}, f"CRC error in {bad}"
            return {name: archive.read(name) for name in archive.namelist()}, None
    except Exception as exc:
        return {}, str(exc)


def count_xml(parts: dict[str, bytes], marker: bytes) -> int:
    return sum(content.count(marker) for name, content in parts.items() if name.endswith(".xml"))


def count_xml_tag(parts: dict[str, bytes], tag: bytes) -> int:
    pattern = re.compile(rb"<w:" + re.escape(tag) + rb"(?:\s|>)")
    return sum(len(pattern.findall(content)) for name, content in parts.items() if name.endswith(".xml"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--expected-title")
    parser.add_argument("--forbidden-term", action="append", default=[])
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    parts, zip_error = read_zip_parts(args.docx)
    result: dict[str, object] = {
        "path": str(args.docx),
        "size": args.docx.stat().st_size if args.docx.exists() else None,
        "zip_error": zip_error,
        "errors": [],
        "warnings": [],
    }
    if zip_error:
        result["errors"].append("DOCX ZIP integrity failed")
    else:
        document = Document(str(args.docx))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        text = "\n".join(paragraphs)
        result.update({
            "paragraphs": len(paragraphs),
            "sections": len(document.sections),
            "tables": len(document.tables),
            "seven_sections": {name: text.count(name) for name in SECTIONS},
            "highlight_count": count_xml(parts, b"<w:highlight"),
            "revision_count": count_xml_tag(parts, b"ins") + count_xml_tag(parts, b"del"),
            "comment_anchor_count": count_xml(parts, b"<w:commentRangeStart"),
            "comments_part": any(name.startswith("word/comments") for name in parts),
            "todo_count": len(re.findall(r"TODO|TBD|待补充|占位符", text, flags=re.I)),
            "long_ascii_space_runs": len(re.findall(r" {3,}", text.replace("目    录", "目录"))),
        })

        fixed_errors, fixed_details = check_fixed_format(document)
        result["fixed_format"] = fixed_details
        result["errors"].extend(fixed_errors)

        missing = [name for name in SECTIONS if name not in text]
        if missing:
            result["errors"].append(f"Missing first-level sections: {missing}")
        if result["highlight_count"]:
            result["errors"].append("Highlight remains")
        if result["revision_count"]:
            result["errors"].append("Tracked revisions remain")
        if result["comment_anchor_count"] or result["comments_part"]:
            result["errors"].append("Comments remain")
        if result["todo_count"]:
            result["errors"].append("TODO/TBD/placeholder text remains")
        if result["long_ascii_space_runs"]:
            result["warnings"].append("Long ASCII-space runs found; inspect rendering")

        if args.expected_title:
            count = text.count(args.expected_title)
            result["expected_title"] = args.expected_title
            result["title_occurrences"] = count
            if count < 2:
                result["errors"].append("Expected project title is not present in both directory/body")

        forbidden = {term: text.count(term) for term in args.forbidden_term if term and term in text}
        result["forbidden_terms"] = forbidden
        if forbidden:
            result["errors"].append(f"Forbidden terms remain: {list(forbidden)}")

    result["passed"] = not result["errors"]
    encoded = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded, encoding="utf-8")
    print(encoded)
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
