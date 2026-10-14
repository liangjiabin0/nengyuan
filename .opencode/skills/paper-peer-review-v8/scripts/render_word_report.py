#!/usr/bin/env python3
from __future__ import annotations

import argparse, json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ZH = {
    'title':'论文审稿意见','classification':'论文分类','overall':'一、总体评价','claims':'二、核心结论—实验证据矩阵',
    'vars':'三、未定义符号审查（最多3项）','availability':'四、数据与代码可获得性审查',
    'repro':'五、独立原则复现','category':'六、分类专项审查','latest':'七、Latest Relevant Literature and Comparator Audit（最重要3项）',
    'theory':'八、Theoretical Basis and Derivation Audit（最重要3项）','major':'九、主要审稿意见','minor':'十、次要意见',
    'verdict':'十一、审稿结论','limits':'十二、审查局限','primary':'Primary Category','secondary':'Secondary Category',
    'decision':'建议结论','confidence':'置信度','id':'编号','claim':'核心结论','evidence':'现有证据','status':'支撑状态','assessment':'判断',
    'symbol':'变量/符号','location':'位置','problem':'问题','impact':'影响','data':'数据状态','code':'代码状态','item':'项目','value':'内容',
    'work':'论文','year':'年份','source':'期刊／会议／预印本','relevance':'相关性','cited':'已引用','compared':'已实证比较',
    'theory_claim':'理论结论','support':'理论支撑状态','assumptions':'关键假设','resolution':'最低解决要求'
}
EN = {
    'title':'Peer Review Report','classification':'Paper Classification','overall':'1. Overall Assessment','claims':'2. Central Claim–Evidence Matrix',
    'vars':'3. Key Undefined Variables/Symbols (Top 3)','availability':'4. Data and Code Availability Audit',
    'repro':'5. Principle-Level Independent Reproduction','category':'6. Category-Specific Audit','latest':'7. Latest Relevant Literature and Comparator Audit (Top 3)',
    'theory':'8. Theoretical Basis and Derivation Audit (Top 3)','major':'9. Major Comments','minor':'10. Minor Comments',
    'verdict':'11. Final Recommendation','limits':'12. Limitations of Review','primary':'Primary Category','secondary':'Secondary Category',
    'decision':'Recommendation','confidence':'Confidence','id':'ID','claim':'Central Claim','evidence':'Evidence','status':'Support Status','assessment':'Assessment',
    'symbol':'Variable/Symbol','location':'Location','problem':'Problem','impact':'Impact','data':'Data Status','code':'Code Status','item':'Item','value':'Content',
    'work':'Work','year':'Year','source':'Journal / Conference / Preprint','relevance':'Relevance','cited':'Cited','compared':'Experimentally Compared',
    'theory_claim':'Theoretical Claim','support':'Support Status','assumptions':'Key Assumptions','resolution':'Minimum Resolution'
}

def clear_body(doc: Document):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)

def set_font(run, size=10.5, bold=None):
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    rpr=run._r.get_or_add_rPr(); rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement('w:rFonts'); rpr.insert(0,rfonts)
    rfonts.set(qn('w:eastAsia'),'宋体'); rfonts.set(qn('w:ascii'),'Times New Roman'); rfonts.set(qn('w:hAnsi'),'Times New Roman')

def add_para(doc, text='', bold=False, size=10.5, align=None):
    p=doc.add_paragraph(); r=p.add_run(str(text)); set_font(r,size,bold)
    p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.15
    if align is not None: p.alignment=align
    return p

def add_heading(doc, text):
    p=doc.add_paragraph(); r=p.add_run(text); set_font(r,12,bold=True); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(4); p.paragraph_format.keep_with_next=True
    return p

def fill_cell(cell,text,bold=False,size=9):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run('' if text is None else str(text)); set_font(r,size,bold); p.paragraph_format.space_after=Pt(0); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    for j,h in enumerate(headers): fill_cell(t.rows[0].cells[j],h,True,8.8)
    for row in rows:
        cells=t.add_row().cells
        for j,v in enumerate(row): fill_cell(cells[j],v,False,8.5)
    return t

def truth(v):
    if v is True: return 'Yes'
    if v is False: return 'No'
    return '' if v is None else str(v)

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--input',required=True,type=Path); ap.add_argument('--output',required=True,type=Path); ap.add_argument('--template',type=Path); ap.add_argument('--ui-language',choices=['zh-CN','en'],required=True)
    a=ap.parse_args(); d=json.loads(a.input.read_text(encoding='utf-8-sig')); ui=ZH if a.ui_language=='zh-CN' else EN
    doc=Document(a.template) if a.template else Document()
    clear_body(doc)
    for sec in doc.sections:
        sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2.2); sec.right_margin=Cm(2.2)
    title=d.get('metadata',{}).get('paper_title','')
    add_para(doc,ui['title'],True,16,WD_ALIGN_PARAGRAPH.CENTER); add_para(doc,title,False,11,WD_ALIGN_PARAGRAPH.CENTER)
    add_heading(doc,ui['classification'])
    c=d.get('classification',{}); table(doc,[ui['item'],ui['value']],[[ui['primary'],c.get('primary_category','')],[ui['secondary'],c.get('secondary_category') or ''],[ui['decision'],d.get('verdict',{}).get('decision','')],[ui['confidence'],d.get('verdict',{}).get('confidence','')]])
    add_heading(doc,ui['overall']); add_para(doc,d.get('summary','')); add_para(doc,d.get('verdict',{}).get('rationale',''))
    add_heading(doc,ui['claims']); rows=[]
    for x in d.get('conclusion_support',{}).get('claims',[]): rows.append([x.get('id',''),x.get('claim',''),'; '.join(map(str,x.get('supporting_evidence',[]))),x.get('support_status',''),x.get('evidence_assessment','')])
    table(doc,[ui['id'],ui['claim'],ui['evidence'],ui['status'],ui['assessment']],rows)
    add_heading(doc,ui['vars']); rows=[]
    for x in d.get('undefined_variables',[])[:3]: rows.append([x.get('symbol',''),x.get('location',''),x.get('problem',''),x.get('impact','')])
    table(doc,[ui['symbol'],ui['location'],ui['problem'],ui['impact']],rows or [['','','','']])
    add_heading(doc,ui['availability']); av=d.get('data_code_availability',{}); table(doc,[ui['item'],ui['value']],[[ui['data'],av.get('data_status','')],[ui['code'],av.get('code_status','')],[ui['evidence'],av.get('evidence','')],['Gaps','; '.join(map(str,av.get('gaps',[])))]])
    add_heading(doc,ui['repro']); rp=d.get('reproduction',{}); add_para(doc, f"Mode: {rp.get('mode','')} | Readiness: {rp.get('readiness','')} | Scope: {rp.get('verification_scope','')}"); add_para(doc,rp.get('data_used',''))
    for s in rp.get('observations',[]): add_para(doc,'• '+str(s))
    rows=[]
    for x in rp.get('barriers',[])[:3]: rows.append([x.get('item',''),x.get('impact',''),x.get('required_action','')])
    table(doc,[ui['item'],ui['impact'],ui['resolution']],rows or [['','','']])
    add_heading(doc,ui['category']); cat=d.get('category_audit',{}); block=cat.get({'data_science':'data_science','optimization_planning_operation':'optimization','system_stability_control':'stability_control'}.get(c.get('primary_category',''),'data_science'))
    add_para(doc, json.dumps(block,ensure_ascii=False,indent=2) if isinstance(block,(dict,list)) else (block or ''))
    add_heading(doc,ui['latest']); rows=[]
    for x in d.get('latest_literature_comparison',{}).get('recent_relevant_works',[])[:3]: rows.append([x.get('title',''),x.get('year',''),x.get('journal') or x.get('source',''),x.get('relevance',''),truth(x.get('cited_in_manuscript')),truth(x.get('experimentally_compared'))])
    table(doc,[ui['work'],ui['year'],ui['source'],ui['relevance'],ui['cited'],ui['compared']],rows or [['','','','','','']])
    add_heading(doc,ui['theory']); rows=[]
    for x in d.get('theoretical_analysis',{}).get('claims',[])[:3]: rows.append([x.get('id',''),x.get('claim',''),x.get('theoretical_support_status',''),'; '.join(map(str,x.get('assumptions',[]))),x.get('assessment',''),x.get('required_resolution','')])
    table(doc,[ui['id'],ui['theory_claim'],ui['support'],ui['assumptions'],ui['assessment'],ui['resolution']],rows or [['','','','','','']])
    add_heading(doc,ui['major'])
    for x in d.get('issues',[]):
        if x.get('severity') in ('critical','major'):
            p=add_para(doc,f"{x.get('id','')}. {x.get('problem','')}",True); p.paragraph_format.keep_with_next=True
            add_para(doc,f"Evidence: {x.get('evidence','')}")
            if x.get('reproduction_observation'): add_para(doc,f"Reproduction Observation: {x.get('reproduction_observation')}")
            add_para(doc,f"Consequence: {x.get('consequence','')}")
            add_para(doc,f"Required Action: {x.get('required_action','')}")
            add_para(doc,f"Acceptance Criteria: {x.get('acceptance_criteria','')}")
    add_heading(doc,ui['minor'])
    for i,x in enumerate(d.get('minor_comments',[]),1): add_para(doc,f"{i}. {x}")
    add_heading(doc,ui['verdict']); add_para(doc,d.get('verdict',{}).get('decision',''),True); add_para(doc,d.get('verdict',{}).get('rationale',''))
    add_heading(doc,ui['limits'])
    for x in d.get('limitations_of_review',[]): add_para(doc,'• '+str(x))
    for sec in doc.sections:
        fp=sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph(); fp.text='Peer review | evidence, reproduction, latest comparator, theory'; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    a.output.parent.mkdir(parents=True,exist_ok=True); doc.save(a.output); print(a.output)
if __name__=='__main__': main()
